"""Build the compact Daily data contract embedded in the pricing dashboard."""
from __future__ import annotations

import csv
import datetime as dt
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from clean_price_data import canonicalize_model


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DAILY_WORKBOOK = ROOT / "Price Check Daily.xlsx"
APPLE_WEEK_ANCHOR_DATE = dt.date(2026, 9, 6)
APPLE_WEEK_ANCHOR_ID = "W11Q4FY26"
PARTNER_NAMES = {
    "FPT": "FPT",
    "VIETTEL": "Viettel",
    "CPS": "CPS",
    "MW": "MW",
    "SHOPDUNK": "Shopdunk",
}


def _parse_datetime(value) -> dt.datetime:
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time())
    parsed = dt.datetime.fromisoformat(str(value))
    return parsed.replace(tzinfo=None) if parsed.tzinfo else parsed


def apple_week_id_for_date(value: dt.date) -> str:
    """Map a calendar date to a 13-week Apple fiscal quarter."""
    delta_weeks = (value - APPLE_WEEK_ANCHOR_DATE).days // 7
    anchor_serial = 26 * 52 + (4 - 1) * 13 + (11 - 1)
    fiscal_year, year_week = divmod(anchor_serial + delta_weeks, 52)
    quarter, quarter_week = divmod(year_week, 13)
    return f"W{quarter_week + 1}Q{quarter + 1}FY{fiscal_year}"


def _normalize_partner(value: str) -> str | None:
    return PARTNER_NAMES.get(str(value or "").strip().upper())


def _price(row: dict) -> int | None:
    for field in ("final_price", "value"):
        raw = str(row.get(field, "") or "").replace(".", "").replace(",", "").strip()
        if raw.isdigit() and int(raw) > 0:
            return int(raw)
    return None


def _logged_runs(workbook_path: Path) -> list[dict]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Run Log" not in workbook.sheetnames:
        return []
    sheet = workbook["Run Log"]
    headers = {cell.value: index for index, cell in enumerate(sheet[1]) if cell.value}
    required = {"Run ID", "Completed At", "Snapshot Path"}
    if not required.issubset(headers):
        return []
    runs = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        run_id = values[headers["Run ID"]]
        completed_at = values[headers["Completed At"]]
        snapshot_path = values[headers["Snapshot Path"]]
        if not run_id or not completed_at or not snapshot_path:
            continue
        path = Path(str(snapshot_path))
        if not path.is_absolute():
            path = workbook_path.parent / path
        if path.is_file():
            runs.append({
                "run_id": str(run_id),
                "completed_at": _parse_datetime(completed_at),
                "snapshot_path": path,
            })
    return sorted(runs, key=lambda item: (item["completed_at"], item["run_id"]))


def _successful_observations(snapshot_path: Path, fallback_time: dt.datetime):
    with snapshot_path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            status = str(row.get("status", "")).upper()
            if status not in {"OK", "OOS"}:
                continue
            partner = _normalize_partner(row.get("retailer", ""))
            model = canonicalize_model(row.get("model", ""))
            if not partner or not model["canonical_model"]:
                continue
            fetched_at = row.get("fetched_at") or fallback_time.isoformat(timespec="seconds")
            try:
                observed_at = _parse_datetime(fetched_at)
            except ValueError:
                observed_at = fallback_time
            yield {
                "key": (model["canonical_model"], partner),
                "category": model["category"],
                "canonical_model": model["canonical_model"],
                "partner": partner,
                "price_vnd": None if status == "OOS" else _price(row),
                "stock_status": "OOS" if status == "OOS" else "In stock",
                "observed_at": observed_at,
                "confidence": str(row.get("confidence", "") or ""),
                "risk_flags": str(row.get("risk_flags", "") or ""),
            }


def load_daily_payload(
    workbook_path: Path = DEFAULT_DAILY_WORKBOOK,
    retention_weeks: int = 13,
) -> dict[str, list]:
    """Return effective end-of-run Daily states for logged snapshot dates."""
    runs = _logged_runs(Path(workbook_path))
    if not runs:
        return {"dailyDates": [], "dailyWeeks": [], "dailyRows": []}

    runs_by_date: dict[dt.date, list[dict]] = defaultdict(list)
    for run in runs:
        runs_by_date[run["completed_at"].date()].append(run)

    latest_date = max(runs_by_date)
    latest_week_start = latest_date - dt.timedelta(days=(latest_date.weekday() + 1) % 7)
    cutoff = latest_week_start - dt.timedelta(weeks=max(retention_weeks - 1, 0))
    state: dict[tuple[str, str], dict] = {}
    output_rows = []
    output_dates = []

    for run_date in sorted(runs_by_date):
        for run in runs_by_date[run_date]:
            for observation in _successful_observations(run["snapshot_path"], run["completed_at"]):
                old = state.get(observation["key"])
                if old is None or observation["observed_at"] >= old["observed_at"]:
                    state[observation["key"]] = observation
        if run_date < cutoff:
            continue
        date_text = run_date.isoformat()
        output_dates.append(date_text)
        week_id = apple_week_id_for_date(run_date)
        for record in sorted(state.values(), key=lambda item: (item["category"], item["canonical_model"], item["partner"])):
            output_rows.append({
                "date": date_text,
                "week_id": week_id,
                "category": record["category"],
                "canonical_model": record["canonical_model"],
                "partner": record["partner"],
                "price_vnd": record["price_vnd"],
                "stock_status": record["stock_status"],
                "observed_at": record["observed_at"].isoformat(timespec="seconds"),
                "stale": record["observed_at"].date() != run_date,
                "confidence": record["confidence"],
                "risk_flags": record["risk_flags"],
            })

    return {
        "dailyDates": output_dates,
        "dailyWeeks": list(dict.fromkeys(apple_week_id_for_date(dt.date.fromisoformat(day)) for day in output_dates)),
        "dailyRows": output_rows,
    }
