#!/usr/bin/env python3
"""
Update weekly prices in Price Check.xlsx from direct product links.

Input contract:
  Link Product.xlsx
    - Column A: Model
    - Other columns: retailer names, e.g. Viettel, MW, FPT, CPS, Shopdunk
    - Cells: direct product URLs to crawl

The tool creates the next weekly sheet by copying the latest weekly sheet,
updates prices for links present in Link Product.xlsx, highlights changed prices
in yellow, and writes comments with prior price and direct discount details.
"""

from __future__ import annotations

import argparse
import copy
import csv
import datetime as dt
import html
import json
import os
import re
import shutil
import ssl
import subprocess
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from dataclasses import dataclass
from http.cookiejar import CookieJar
from pathlib import Path
from typing import Iterable
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import quote_sheetname

try:
    import certifi
except ImportError:  # pragma: no cover - optional runtime dependency.
    certifi = None

try:
    import truststore
except ImportError:  # pragma: no cover - optional runtime dependency.
    truststore = None


DEFAULT_WORKBOOK = Path("Price Check.xlsx")
DEFAULT_LINK_WORKBOOK = Path("Link Product.xlsx")
DEFAULT_OVERRIDE_FILE = Path("Price Overrides.csv")
DEFAULT_REPORT_DIR = Path("outputs/reports")
DEFAULT_BACKUP_DIR = Path("outputs/backups")
PRICE_CHANGED_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")
PRICE_REVIEW_FILL = PatternFill(fill_type="solid", fgColor="F4B183")
REVIEW_STATUS_FLAGS = {
    "WARN_HTML_RENDER_MISMATCH",
    "WARN_RENDER_UNAVAILABLE",
    "WARN_STOCK_ACTION_UNVERIFIED",
    "WARN_OUTLIER",
    "WARN_SUSPICIOUS_LOW",
    "WARN_FETCH_STATUS",
}

RETAILER_ALIASES = {
    "cellphones": "CPS",
    "cellphone s": "CPS",
    "cps": "CPS",
    "fpt": "FPT",
    "fptshop": "FPT",
    "mw": "MW",
    "tgdd": "MW",
    "thegioididong": "MW",
    "topzone": "TZ",
    "tz": "TZ",
    "viettel": "Viettel",
    "viettelstore": "Viettel",
    "shopdunk": "Shopdunk",
    "shop dunk": "Shopdunk",
    "sd": "Shopdunk",
    "hoang ha": "Hoang Ha",
    "hoàng hà": "Hoang Ha",
    "hh": "Hoang Ha",
    "shopee": "Shopee",
}

VALID_RETAILERS = {"TZ", "MW", "FPT", "Viettel", "CPS", "Hoang Ha", "Shopdunk", "Shopee"}
OOS_PHRASES = (
    "tam het hang",
    "khong co hang",
    "hang sap ve",
    "sap ve",
    "het hang",
    "ngung kinh doanh",
    "lien he tu van",
    "sold out",
    "out of stock",
)
PURCHASE_ACTION_OOS_PHRASES = (
    "cho lien he",
    "lien he tu van",
    "sap ve hang",
    "san pham ngung kinh doanh",
    "tam het hang",
    "khong co hang",
    "hang sap ve",
    "ngung kinh doanh",
    "sold out",
    "out of stock",
)
PURCHASE_ACTION_OK_PHRASES = (
    "mua hang",
    "dat hang",
    "mua ngay",
    "them vao gio",
)
TRUSTED_RENDER_METHOD_PREFIXES = (
    "rendered_dom:.st-price-main",
    "rendered_dom:.price-main",
    "rendered_dom:[data-testid*='price']",
    "rendered_dom:.product-detail .price",
    "rendered_dom:.box_saving .bs_price strong",
    "rendered_dom:.box_saving",
    "rendered_dom:.box-price-present",
    "rendered_dom:.price-one",
    "rendered_dom:.sale-price",
    "rendered_dom:.box-product-price",
    "rendered_dom:[id^='price-value-']",
    "rendered_dom:.new-price",
)
RENDER_BATCH_SIZE = 40


@dataclass
class Target:
    model: str
    retailer: str
    url: str


@dataclass
class Result:
    model: str
    retailer: str
    url: str
    value: int | str | None
    status: str
    error: str = ""
    fetched_at: str = ""
    base_value: int | str | None = None
    discount: int = 0
    promo_note: str = ""
    source_method: str = ""
    html_price: int | str | None = None
    rendered_price: int | None = None
    previous_week_price: int | str | None = None
    risk_flags: str = ""
    confidence: str = ""
    decision_note: str = ""
    availability: str = ""
    availability_method: str = ""
    availability_text: str = ""
    purchase_action: str = ""
    purchase_action_method: str = ""
    purchase_action_text: str = ""


@dataclass(frozen=True)
class BrowserRuntime:
    node_bin: Path | None
    node_path: Path | None
    chromium_bin: Path | None


def resolve_browser_runtime(
    env: dict[str, str] | None = None,
    platform_name: str | None = None,
    which=shutil.which,
) -> BrowserRuntime:
    """Resolve browser tooling from explicit overrides, then platform defaults."""
    values = os.environ if env is None else env
    platform_name = platform_name or sys.platform
    mac_node = Path("/Users/vutrungnghia/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node")
    mac_modules = Path("/Users/vutrungnghia/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules")
    mac_chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")

    def configured(name: str) -> Path | None:
        value = values.get(name, "").strip()
        return Path(value) if value else None

    node_bin = configured("PRICE_CHECK_NODE_BIN")
    node_path = configured("PRICE_CHECK_NODE_PATH")
    chromium_bin = configured("PRICE_CHECK_CHROMIUM_BIN")
    if platform_name == "darwin":
        node_bin = node_bin or mac_node
        node_path = node_path or mac_modules
        chromium_bin = chromium_bin or mac_chrome
    else:
        node_bin = node_bin or next((Path(path) for name in ("node", "nodejs") if (path := which(name))), None)
        chromium_bin = chromium_bin or next(
            (
                Path(path)
                for name in ("chromium", "chromium-browser", "google-chrome", "google-chrome-stable")
                if (path := which(name))
            ),
            None,
        )
    return BrowserRuntime(node_bin=node_bin, node_path=node_path, chromium_bin=chromium_bin)


def normalize_name(value: str) -> str:
    value = " ".join((value or "").strip().lower().split())
    return RETAILER_ALIASES.get(value, value.title())


def workbook_sheets(workbook_path: Path) -> list[str]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    return wb.sheetnames


def latest_week_sheet(sheet_names: Iterable[str]) -> str:
    candidates = [s for s in sheet_names if re.match(r"^W\d+Q\d+FY\d+", s.strip())]
    if not candidates:
        raise ValueError("No weekly sheet name like W7Q3FY26 was found.")
    return candidates[-1]


def next_week_name(sheet_name: str) -> str:
    match = re.match(r"^W(\d+)Q(\d+)FY(\d+)", sheet_name.strip())
    if not match:
        raise ValueError(f"Cannot infer next week from sheet name: {sheet_name}")
    week, quarter, fy = map(int, match.groups())
    week += 1
    if week > 13:
        week = 1
        quarter += 1
        if quarter > 4:
            quarter = 1
            fy += 1
    return f"W{week}Q{quarter}FY{fy:02d}"


def read_link_targets(link_workbook: Path, retailer_filter: set[str] | None = None) -> list[Target]:
    wb = load_workbook(link_workbook, read_only=True, data_only=False)
    ws = wb.active
    headers = []
    for col in range(2, ws.max_column + 1):
        retailer = normalize_name(str(ws.cell(1, col).value or ""))
        if retailer in VALID_RETAILERS and (retailer_filter is None or retailer in retailer_filter):
            headers.append((col, retailer))

    targets: list[Target] = []
    for row in range(2, ws.max_row + 1):
        model = str(ws.cell(row, 1).value or "").strip()
        if not model:
            continue
        for col, retailer in headers:
            url = str(ws.cell(row, col).value or "").strip()
            if url.startswith("http"):
                targets.append(Target(model=model, retailer=retailer, url=url))
    return targets


def read_link_retailers(link_workbook: Path) -> set[str]:
    wb = load_workbook(link_workbook, read_only=True, data_only=False)
    ws = wb.active
    return {
        retailer
        for col in range(2, ws.max_column + 1)
        if (retailer := normalize_name(str(ws.cell(1, col).value or ""))) in VALID_RETAILERS
    }


def missing_link_results(
    workbook_path: Path,
    source_sheet: str,
    link_workbook: Path,
    all_targets: list[Target],
) -> list[Result]:
    linked = {(target.model, normalize_name(target.retailer)) for target in all_targets}
    configured_retailers = read_link_retailers(link_workbook)
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    ws = wb[source_sheet]
    workbook_retailers = {
        normalize_name(str(ws.cell(1, col).value or ""))
        for col in range(2, ws.max_column + 1)
    }
    results: list[Result] = []
    for row in range(2, ws.max_row + 1):
        model = str(ws.cell(row, 1).value or "").strip()
        if not model:
            continue
        for retailer in configured_retailers & workbook_retailers:
            if (model, retailer) not in linked:
                results.append(
                    Result(
                        model=model,
                        retailer=retailer,
                        url="",
                        value="OOS",
                        status="OOS",
                        source_method="missing_direct_link",
                        confidence="OOS_CONFIRMED",
                        decision_note="no direct product URL in Link Product.xlsx",
                        availability="OOS",
                        availability_method="missing_direct_link",
                        availability_text="No direct product URL in Link Product.xlsx",
                    )
                )
    return results


def read_price_overrides(override_file: Path) -> dict[str, tuple[int, str]]:
    if not override_file.exists():
        return {}
    overrides: dict[str, tuple[int, str]] = {}
    with override_file.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            url = str(row.get("url") or "").strip()
            value = parse_price_number(str(row.get("value") or ""))
            note = str(row.get("note") or "").strip()
            if url.startswith("http") and value:
                overrides[url] = (value, note)
    return overrides


def ssl_context() -> ssl.SSLContext:
    # The bundled Python runtime does not inherit macOS Keychain roots. Read
    # the same system and user keychains used by macOS rather than relying on
    # an optional package that may not persist between runtime sessions.
    if sys.platform == "darwin":
        keychains = (
            "/System/Library/Keychains/SystemRootCertificates.keychain",
            "/Library/Keychains/System.keychain",
            str(Path.home() / "Library/Keychains/login.keychain-db"),
        )
        pem_parts: list[str] = []
        for keychain in keychains:
            try:
                completed = subprocess.run(
                    ["security", "find-certificate", "-a", "-p", keychain],
                    capture_output=True,
                    text=True,
                    timeout=10,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired):
                continue
            if completed.returncode == 0 and "BEGIN CERTIFICATE" in completed.stdout:
                pem_parts.append(completed.stdout)
        if pem_parts:
            return ssl.create_default_context(cadata="\n".join(pem_parts))
    if certifi is not None:
        return ssl.create_default_context(cafile=certifi.where())
    # The bundled Python runtime does not inherit macOS Keychain roots.  Its
    # system CA bundle still verifies certificates without weakening TLS.
    system_ca_bundle = Path("/etc/ssl/cert.pem")
    if system_ca_bundle.is_file():
        return ssl.create_default_context(cafile=str(system_ca_bundle))
    return ssl.create_default_context()


def request_html(url: str, timeout: int = 30) -> str:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"
        ),
        "Accept-Language": "vi-VN,vi;q=0.9,en-US;q=0.8,en;q=0.7",
    }
    if "cellphones.com.vn" in url:
        headers["Cookie"] = "cps_province_id=24; cps_region=1"
    request = Request(
        url,
        headers=headers,
    )
    context = ssl_context()
    if "cellphones.com.vn" not in url:
        with urlopen(request, timeout=timeout, context=context) as response:
            charset = response.headers.get_content_charset() or "utf-8"
            return response.read().decode(charset, errors="replace")

    best_html = ""
    best_price: int | None = None
    last_html = ""
    for _ in range(5):
        with urlopen(request, timeout=timeout, context=context) as response:
            charset = response.headers.get_content_charset() or "utf-8"
            raw_html = response.read().decode(charset, errors="replace")
        last_html = raw_html
        sale_price = re.search(r'<div class=["\']sale-price["\']>\s*([^<]+)', raw_html, flags=re.I)
        price = parse_price_number(html.unescape(sale_price.group(1))) if sale_price else None
        if best_price is None or (price is not None and price < best_price):
            best_price = price
            best_html = raw_html
    return best_html or last_html


def shopdunk_stock_probe(url: str) -> tuple[str, str, str] | None:
    """Resolve Shopdunk's selected SKU and final stock gate used by its UI."""
    if "shopdunk.com" not in url.lower():
        return None
    jar = CookieJar()
    opener = urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(jar),
        urllib.request.HTTPSHandler(context=ssl_context()),
    )
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"
        ),
        "Accept-Language": "vi-VN,vi;q=0.9,en-US;q=0.8,en;q=0.7",
        "Accept": "text/html,application/xhtml+xml",
    }
    try:
        with opener.open(urllib.request.Request(url, headers=headers), timeout=30) as response:
            raw_html = response.read().decode(response.headers.get_content_charset() or "utf-8", errors="replace")
    except (HTTPError, URLError, TimeoutError) as exc:
        return "UNVERIFIED", "shopdunk_stock_probe_error", compact_evidence(str(exc))

    form_match = re.search(
        r'(?is)<form[^>]+id\s*=\s*(?:["\']product-details-form["\']|product-details-form)[^>]*>(.*?)</form>',
        raw_html,
    )
    product_match = re.search(r"(?:data-promotion-url|productId=|/details/)(?:[^\"']*?productId=)?(\d+)", raw_html, flags=re.I)
    if not form_match or not product_match:
        return "UNVERIFIED", "shopdunk_stock_probe_missing_form", "Selected SKU form was not found"
    product_id = product_match.group(1)
    form_html = form_match.group(1)
    pairs: list[tuple[str, str]] = []
    def attribute_value(attrs: str, name: str) -> str:
        match = re.search(rf"\b{name}\s*=\s*(?:[\"']([^\"']*)[\"']|([^\s>]+))", attrs, flags=re.I)
        return html.unescape(next((value for value in match.groups() if value is not None), "")) if match else ""

    for match in re.finditer(r"(?is)<input\b([^>]*)>", form_html):
        attrs = match.group(1)
        name = attribute_value(attrs, "name")
        value = attribute_value(attrs, "value")
        if not name:
            continue
        input_type = (attribute_value(attrs, "type") or "text").lower()
        checked = bool(re.search(r"\bchecked(?:\s*=\s*[\"'][^\"']*[\"'])?", attrs, flags=re.I))
        if input_type in {"radio", "checkbox"} and not checked:
            continue
        if name.startswith("addtocart_"):
            continue
        pairs.append((name, value))
    for match in re.finditer(r"(?is)<select\b([^>]*)>(.*?)</select>", form_html):
        attrs, body = match.groups()
        name = attribute_value(attrs, "name")
        option_match = re.search(r"(?is)<option\b[^>]*selected[^>]*", body)
        if name and option_match:
            value = attribute_value(option_match.group(0), "value")
            if value:
                pairs.append((name, value))
    if not pairs:
        return "UNVERIFIED", "shopdunk_stock_probe_missing_attributes", "Selected variant attributes were not found"

    endpoint = urllib.parse.urljoin(url, f"/shoppingcart/productdetails_attributechange?productId={product_id}&validateAttributeConditions=False&loadPicture=True")
    ajax_headers = {
        **headers,
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "X-Requested-With": "XMLHttpRequest",
        "Referer": url,
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
    }
    try:
        request = urllib.request.Request(
            endpoint,
            data=urllib.parse.urlencode(pairs).encode(),
            headers=ajax_headers,
            method="POST",
        )
        with opener.open(request, timeout=30) as response:
            payload = json.loads(response.read().decode("utf-8", errors="replace"))
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as exc:
        return "UNVERIFIED", "shopdunk_stock_probe_error", compact_evidence(str(exc))

    sku = str(payload.get("sku") or "").strip()
    stock_label = str(payload.get("stockAvailability") or "").strip()
    if not sku:
        return "UNVERIFIED", "shopdunk_stock_probe_missing_sku", compact_evidence(stock_label or "SKU was not returned")
    stock_endpoint = urllib.parse.urljoin(url, "/bkc/CheckProductAvailablity")
    try:
        request = urllib.request.Request(
            stock_endpoint,
            data=urllib.parse.urlencode({"sku": sku}).encode(),
            headers=ajax_headers,
            method="POST",
        )
        with opener.open(request, timeout=30) as response:
            final_gate = response.read().decode("utf-8", errors="replace").strip().lower()
    except (HTTPError, URLError, TimeoutError) as exc:
        return "UNVERIFIED", "shopdunk_stock_probe_error", compact_evidence(str(exc))
    evidence = f"sku={sku}; stockAvailability={stock_label or 'unknown'}; finalGate={final_gate or 'unknown'}"
    if final_gate == "true":
        return "IN_STOCK", "shopdunk_stock_endpoint", evidence
    if final_gate == "false":
        return "OOS", "shopdunk_stock_endpoint", evidence
    return "UNVERIFIED", "shopdunk_stock_probe_unknown", evidence


def apply_shopdunk_stock_probe(result: Result, target: Target) -> Result:
    probe = shopdunk_stock_probe(target.url)
    if not probe:
        return result
    status, method, evidence = probe
    result.purchase_action = status
    result.purchase_action_method = method
    result.purchase_action_text = evidence
    if status == "OOS":
        result.value = "OOS"
        result.base_value = "OOS"
        result.status = "OOS"
        result.error = ""
        result.availability = "OOS"
        result.availability_method = method
        result.availability_text = evidence
        result.confidence = "OOS_CONFIRMED"
        result.decision_note = f"confirmed unavailable via {method}"
    elif status == "IN_STOCK" and result.status == "OOS":
        recovered_price = result.html_price if isinstance(result.html_price, int) else None
        if recovered_price:
            result.value = recovered_price
            result.base_value = recovered_price
            result.status = "OK"
            result.error = ""
            result.availability = ""
            result.availability_method = ""
            result.availability_text = ""
            result.confidence = "VERIFIED_RENDER"
            result.decision_note = f"Shopdunk stock endpoint confirmed purchasable via {method}"
    return result


def strip_tags(raw: str) -> str:
    text = re.sub(r"(?is)<script.*?</script>|<style.*?</style>", " ", raw)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return html.unescape(re.sub(r"\s+", " ", text)).strip()


def parse_price_number(raw: str) -> int | None:
    if not raw:
        return None
    compact = raw.replace("\xa0", " ")
    if re.search(r"\b(oos|n/a|hết|het|out of stock)\b", compact, flags=re.I):
        return None
    decimal_integer = re.search(r"(?<![\d.,])(\d{6,9})[.,]0{1,2}(?!\d)", compact)
    if decimal_integer:
        return int(decimal_integer.group(1))
    digits = re.sub(r"\D", "", compact)
    if not digits:
        return None
    number = int(digits)
    if number < 1000:
        return None
    return number


def normalize_search_text(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value or "")
    ascii_text = "".join(char for char in decomposed if unicodedata.category(char) != "Mn")
    return " ".join(ascii_text.lower().split())


def find_oos_phrase(value: str) -> str | None:
    normalized = normalize_search_text(value)
    return next((phrase for phrase in OOS_PHRASES if phrase in normalized), None)


def extract_purchase_action(raw_html: str) -> tuple[str, str, str] | None:
    """Find an explicitly visible purchase CTA in server-rendered product markup."""
    candidates: list[tuple[str, str, str]] = []
    hidden_oos_text: list[str] = []
    tag_pattern = re.compile(r"(?is)<(button|a)\b([^>]*)>(.*?)</\1>")
    for match in tag_pattern.finditer(raw_html):
        tag, attrs, body = match.groups()
        attrs_lower = attrs.lower()
        text = strip_tags(body)
        normalized = normalize_search_text(text)
        if not text:
            continue
        parent_window = raw_html[max(0, match.start() - 600) : match.start()]
        parent_open = parent_window.rfind("<div")
        parent_close = parent_window.rfind("</div")
        parent_attrs = parent_window[parent_open : parent_window.find(">", parent_open) + 1] if parent_open > parent_close else ""
        is_hidden = bool(
            re.search(
                r"(?:d-none|hidden|display\s*:\s*none|visibility\s*:\s*hidden)",
                f"{attrs_lower} {parent_attrs.lower()}",
            )
        )
        if is_hidden:
            if any(phrase in normalized for phrase in PURCHASE_ACTION_OOS_PHRASES):
                hidden_oos_text.append(text)
            continue
        if tag == "a" and re.search(r"href\s*=\s*[\"']tel:", attrs_lower):
            continue
        action_attr = f"{attrs_lower} {normalized}"
        is_action_element = tag == "button" or re.search(
            r"(?:add[-_]to[-_]cart|buy|order|checkout|cart|product[-_]detail|btn[-_]buy|btn[-_]order|btndattruoc|dathang)",
            action_attr,
        )
        is_action_text = any(phrase in normalized for phrase in PURCHASE_ACTION_OK_PHRASES + PURCHASE_ACTION_OOS_PHRASES)
        if not is_action_element and not is_action_text:
            continue
        if any(phrase in normalized for phrase in PURCHASE_ACTION_OOS_PHRASES):
            candidates.append(("OOS", f"html_action:{tag}", text))
        elif any(phrase in normalized for phrase in PURCHASE_ACTION_OK_PHRASES):
            candidates.append(("IN_STOCK", f"html_action:{tag}", text))
    if hidden_oos_text and any(item[0] == "IN_STOCK" for item in candidates):
        return (
            "UNVERIFIED",
            "html_action_conflict",
            f"visible {candidates[0][2]}; hidden OOS CTA {hidden_oos_text[0]}",
        )
    if not candidates:
        return None
    # A visible negative CTA is authoritative if a page exposes conflicting controls.
    return next((item for item in candidates if item[0] == "OOS"), candidates[0])


def compact_evidence(value: str, limit: int = 260) -> str:
    compact = " ".join((value or "").split())
    return compact if len(compact) <= limit else f"{compact[: limit - 3]}..."


def canonical_product_url(value: str) -> str:
    return (value or "").split("?", 1)[0].rstrip("/").lower()


def structured_oos(raw_html: str, target_url: str) -> tuple[str, str] | None:
    expected_url = canonical_product_url(target_url)
    for match in re.finditer(
        r'(?is)<script[^>]+type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
        raw_html,
    ):
        try:
            payload = json.loads(html.unescape(match.group(1)).strip())
        except json.JSONDecodeError:
            continue
        stack = payload if isinstance(payload, list) else [payload]
        while stack:
            item = stack.pop()
            if isinstance(item, list):
                stack.extend(item)
                continue
            if not isinstance(item, dict):
                continue
            stack.extend(value for value in item.values() if isinstance(value, (dict, list)))
            offers = item.get("offers")
            offer_items = offers if isinstance(offers, list) else [offers]
            for offer in offer_items:
                if not isinstance(offer, dict):
                    continue
                offer_url = canonical_product_url(str(offer.get("url") or item.get("url") or ""))
                availability = str(offer.get("availability") or "").rsplit("/", 1)[-1]
                if availability.lower() not in {"outofstock", "discontinued", "preorder", "presale", "backorder"}:
                    continue
                if offer_url and offer_url != expected_url:
                    continue
                return "structured_offer_availability", compact_evidence(f"{availability}; {offer_url or expected_url}")

    match = re.search(
        r'<[^>]+itemprop=["\']availability["\'][^>]+(?:content|href)=["\'](?:https?://schema\.org/)?(OutOfStock|Discontinued|PreOrder|PreSale|BackOrder)["\']',
        raw_html,
        flags=re.I,
    )
    if match:
        return "structured_itemprop_availability", match.group(1)
    return None


def price_status_windows(visible: str, price: int | None) -> list[str]:
    if not price:
        return []
    windows = []
    for match in re.finditer(r"([0-9]{1,3}(?:[.,][0-9]{3}){2,})\s*(?:đ|₫|vnd)?", visible, flags=re.I):
        if parse_price_number(match.group(1)) != price:
            continue
        windows.append(visible[max(0, match.start() - 220) : match.end() + 520])
    return windows


def detect_product_oos(
    raw_html: str, visible: str, retailer: str, price: int | None, target_url: str
) -> tuple[str, str] | None:
    normalized_retailer = normalize_name(retailer)
    structured = structured_oos(raw_html, target_url)
    if structured:
        method, evidence = structured
        return method, compact_evidence(evidence)

    action = extract_purchase_action(raw_html)
    if action and action[0] == "OOS":
        return action[1], compact_evidence(action[2])

    # MW product pages contain stock phrases in customer reviews, promotion copy,
    # and recently viewed products.  A nearby price is not proof that the primary
    # SKU is unavailable, so only use its explicit product-status element below.
    if normalized_retailer not in {"MW", "Shopdunk"}:
        for window in price_status_windows(visible, price):
            phrase = find_oos_phrase(window)
            if phrase:
                return f"{normalized_retailer.lower()}_price_context", compact_evidence(window)

    # FPT uses a dedicated noPrice block for discontinued or consultation-only
    # products.  The page can still contain historical/variant prices elsewhere,
    # so this product-detail status must win before accepting any price.
    if normalized_retailer == "FPT":
        fpt_status_blocks = re.finditer(
            r'(?is)<[^>]+(?:id|class)=["\'][^"\']*(?:noPrice|no-price|contact-price)[^"\']*["\'][^>]*>(.{0,1600})',
            raw_html,
        )
        for match in fpt_status_blocks:
            evidence = strip_tags(match.group(1))
            phrase = find_oos_phrase(evidence)
            if phrase:
                return "fpt_product_status_element", compact_evidence(evidence)

    status_blocks = re.finditer(
        r'(?is)<[^>]+(?:class|id)=["\'][^"\']*(?:stock|availability|product-status|productstatus)[^"\']*["\'][^>]*>(.{0,700})',
        raw_html,
    )
    for match in status_blocks:
        evidence = strip_tags(match.group(1))
        phrase = find_oos_phrase(evidence)
        if phrase:
            return "product_status_element", compact_evidence(evidence)
    return None


def oos_result(target: Target, fetched_at: str, method: str, evidence: str) -> Result:
    return Result(
        target.model,
        target.retailer,
        target.url,
        "OOS",
        "OOS",
        fetched_at=fetched_at,
        base_value="OOS",
        source_method=method,
        availability="OOS",
        availability_method=method,
        availability_text=evidence,
        confidence="OOS_CONFIRMED",
        decision_note=f"confirmed unavailable via {method}",
    )


def format_price(value: int | str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def eligible_viettel_discount(model: str, visible_text: str) -> tuple[int, str]:
    """Return broadly accessible Viettel promotions for the selected SKU only."""
    if not visible_text:
        return 0, ""
    text = " ".join(visible_text.split())
    model_lower = normalize_search_text(model)
    eligible_old_customer_model = model_lower.startswith("iphone 17 pro") or model_lower.startswith("iphone air")
    eligible_student_model = "ipad" in model_lower
    matches: list[tuple[int, str]] = []
    saw_excluded_candidate = False
    pattern = re.compile(
        r"(giảm\s+thêm\s+([0-9]{1,3}(?:[.,][0-9]{3})+)\s*(?:đ|d|vnd)?\s+cho\s+.{0,180})",
        flags=re.I,
    )
    for match in pattern.finditer(text):
        promo_text = match.group(1).strip()
        promo_lower = normalize_search_text(promo_text)
        if "tammi" in promo_lower or "goi cuoc" in promo_lower:
            saw_excluded_candidate = True
            continue
        amount = parse_price_number(match.group(2))
        if not amount:
            continue
        if eligible_old_customer_model and "khach hang cu" in promo_lower:
            promo_text = f"Giảm thêm {format_price(amount)} cho khách hàng cũ của Viettel"
        elif eligible_student_model and any(phrase in promo_lower for phrase in ("hoc sinh", "sinh vien", "giao vien")):
            promo_text = f"Giảm thêm {format_price(amount)} cho học sinh, sinh viên, giáo viên"
        else:
            continue
        matches.append((amount, promo_text))
    if not matches:
        if eligible_old_customer_model and not saw_excluded_candidate:
            return 500_000, "Giảm thêm 500,000 cho khách hàng cũ của Viettel"
        return 0, ""
    return max(matches, key=lambda item: item[0])


def eligible_mw_discount(visible_text: str) -> tuple[int, str]:
    """Return directly deducted MW cash promo shown in the main promotion block."""
    if not visible_text:
        return 0, ""
    text = " ".join(visible_text.split())
    marker = "Chọn 1 trong các khuyến mãi"
    marker_index = text.find(marker)
    if marker_index == -1:
        return 0, ""
    window = text[marker_index : marker_index + 500]
    match = re.search(r"Giảm giá\s+([0-9]{1,3}(?:[.,][0-9]{3})+)\s*đ", window, flags=re.I)
    if not match:
        return 0, ""
    amount = parse_price_number(match.group(1))
    if not amount:
        return 0, ""
    return amount, f"Giảm giá {format_price(amount)}"


def eligible_direct_discount(retailer: str, model: str, visible_text: str) -> tuple[int, str]:
    if retailer == "Viettel":
        return eligible_viettel_discount(model, visible_text)
    if retailer == "MW":
        return eligible_mw_discount(visible_text)
    return 0, ""


def extract_jsonld_prices(raw_html: str) -> list[int]:
    prices = []
    for match in re.finditer(
        r'(?is)<script[^>]+type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
        raw_html,
    ):
        payload = html.unescape(match.group(1)).strip()
        try:
            data = json.loads(payload)
        except json.JSONDecodeError:
            continue
        stack = data if isinstance(data, list) else [data]
        while stack:
            item = stack.pop()
            if isinstance(item, dict):
                for key in ("price", "lowPrice", "highPrice"):
                    price = parse_price_number(str(item.get(key) or ""))
                    if price:
                        prices.append(price)
                stack.extend(v for v in item.values() if isinstance(v, (dict, list)))
            elif isinstance(item, list):
                stack.extend(item)
    return prices


def extract_viettel_product_offer(raw_html: str, expected_url: str) -> tuple[int | None, str]:
    """Read the JSON-LD offer that belongs to this exact Viettel product URL."""
    normalized_url = expected_url.rstrip("/")
    for match in re.finditer(
        r'(?is)<script[^>]+type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
        raw_html,
    ):
        try:
            payload = json.loads(html.unescape(match.group(1)).strip())
        except json.JSONDecodeError:
            continue
        stack = payload if isinstance(payload, list) else [payload]
        while stack:
            item = stack.pop()
            if isinstance(item, list):
                stack.extend(item)
                continue
            if not isinstance(item, dict):
                continue
            stack.extend(value for value in item.values() if isinstance(value, (dict, list)))
            offer = item.get("offers")
            offers = offer if isinstance(offer, list) else [offer]
            for candidate in offers:
                if not isinstance(candidate, dict):
                    continue
                offer_url = str(candidate.get("url") or "").rstrip("/")
                if offer_url != normalized_url:
                    continue
                price = parse_price_number(str(candidate.get("price") or ""))
                if price and 500_000 <= price <= 100_000_000:
                    return price, "viettel_exact_jsonld_offer"
    return None, ""


def extract_structured_price(raw_html: str) -> tuple[int | None, str]:
    patterns = (
        (r'"offers"\s*:\s*\{[\s\S]{0,700}?"price"\s*:\s*"?([0-9][0-9.,]*)"?', "structured_offer_price"),
        (r'property=["\']product:price:amount["\'][^>]+content=["\']([^"\']+)["\']', "structured_product_price"),
        (r'<meta[^>]+itemprop=["\']price["\'][^>]+content=["\']([^"\']+)["\']', "structured_itemprop_price"),
        (r'name=["\']price["\'][^>]+content=["\']([^"\']+)["\']', "structured_meta_price"),
        (r'"price"\s*:\s*"?([0-9][0-9.,]*)"?', "structured_json_price"),
    )
    for pattern, method in patterns:
        match = re.search(pattern, raw_html, flags=re.I)
        if match:
            price = parse_price_number(match.group(1))
            if price and 500_000 <= price <= 100_000_000:
                return price, method
    jsonld_prices = extract_jsonld_prices(raw_html)
    if jsonld_prices:
        return min(jsonld_prices), "jsonld_min_price"
    return None, ""


def first_visible_price_after(visible: str, markers: tuple[str, ...]) -> int | None:
    lower = visible.lower()
    start = -1
    for marker in markers:
        index = lower.find(marker.lower())
        if index != -1 and (start == -1 or index < start):
            start = index
    if start == -1:
        return None
    window = visible[start : start + 900]
    for match in re.finditer(r"([0-9]{1,3}(?:[.,][0-9]{3}){2,})\s*(?:đ|₫|vnd)", window, flags=re.I):
        price = parse_price_number(match.group(1))
        if price and 1_000_000 <= price <= 100_000_000:
            return price
    return None


def extract_shopdunk_price(raw_html: str) -> tuple[int | None, str]:
    quoted = r'(?:"([^"]+)"|\'([^\']+)\'|([^\s>]+))'
    patterns = (
        rf'<span[^>]+\bid={quoted}[^>]*>\s*([^<]+)',
        rf'<span[^>]+\bclass={quoted}[^>]*\bnew-price\b[^>]*>\s*([^<]+)',
        rf'<meta[^>]+\bitemprop={quoted}[^>]+\bcontent={quoted}',
    )
    for pattern in patterns:
        for match in re.finditer(pattern, raw_html, flags=re.I):
            groups = [group for group in match.groups() if group]
            if not groups:
                continue
            if pattern.startswith("<span") and not any("price-value" in group or "new-price" in group for group in groups[:-1]):
                continue
            if pattern.startswith("<meta") and not any(group in {"price", "lowPrice"} for group in groups):
                continue
            price = parse_price_number(html.unescape(groups[-1]))
            if price and 500_000 <= price <= 100_000_000:
                return price, "shopdunk_price_value"

    visible = strip_tags(raw_html)
    vat_index = visible.find("(Đã bao gồm VAT)")
    if vat_index != -1:
        before_vat = visible[max(0, vat_index - 250) : vat_index]
        prices = [
            parse_price_number(match.group(1))
            for match in re.finditer(r"([0-9]{1,3}(?:[.,][0-9]{3}){2,})\s*(?:đ|₫|vnd)", before_vat, flags=re.I)
        ]
        prices = [price for price in prices if price and 500_000 <= price <= 100_000_000]
        if prices:
            return prices[0], "shopdunk_vat_window"
    return None, ""


def extract_cps_price(raw_html: str) -> tuple[int | None, str]:
    for pattern, method in (
        (r'<div class=["\']sale-price["\']>\s*([^<]+)', "cps_sale_price"),
        (r'"offers"\s*:\s*\{[\s\S]{0,500}?"price"\s*:\s*"?([0-9][0-9.,]*)"?', "cps_schema_offer"),
        (r'<meta[^>]+itemprop=["\']price["\'][^>]+content=["\']([^"\']+)["\']', "cps_meta_price"),
        (r'"price"\s*:\s*"?([0-9]{7,8})"?', "cps_json_price"),
    ):
        match = re.search(pattern, raw_html, flags=re.I)
        if match:
            price = parse_price_number(match.group(1))
            if price and 500_000 <= price <= 100_000_000:
                return price, method
    return None, ""


def product_capacity(model: str) -> str:
    match = re.search(r"\b([0-9]+)\s*(GB|TB)\b", model, flags=re.I)
    if not match:
        return ""
    return f"{match.group(1)} {match.group(2).upper()}"


def extract_fpt_variant_price(raw_html: str, model: str) -> tuple[int | None, str]:
    capacity = product_capacity(model)
    if not capacity:
        return None, ""
    decoded = html.unescape(raw_html).replace('\\"', '"')
    selected_match = re.search(r'"skuDisplayName"\s*:\s*"([^"]+)"', decoded, flags=re.I)
    selected_capacity = product_capacity(selected_match.group(1)) if selected_match else ""
    if selected_capacity == capacity:
        return None, ""
    section_start = decoded.find('"variants":[{"displayName":"Dung lượng"')
    if section_start == -1:
        return None, ""
    section = decoded[section_start : section_start + 80_000]
    compact_capacity = capacity.replace(" ", "")
    variant_pattern = (
        r'(?:displayName|skuDisplayName|name)":"[^"]*?(?:%s)[^"]*?"'
        r'[\s\S]{0,900}?"(?:finalPrice|currentPrice|price)"\s*:\s*([0-9]{6,9})'
    )
    candidates = []
    for capacity_token in (re.escape(capacity), re.escape(compact_capacity)):
        for match in re.finditer(variant_pattern % capacity_token, section, flags=re.I):
            price = parse_price_number(match.group(1))
            if price and 500_000 <= price <= 100_000_000:
                candidates.append(price)
    if candidates:
        return min(candidates), "fpt_variant_price"
    return None, ""


def extract_site_display_price(raw_html: str, visible: str, retailer: str, model: str = "") -> tuple[int | None, str]:
    if retailer == "MW":
        return first_visible_price_after(visible, ("Online Giá Rẻ Quá", "Giá tại")), "mw_visible_price"
    if retailer == "FPT":
        price, method = extract_fpt_variant_price(raw_html, model)
        if price:
            return price, method
        price, method = extract_structured_price(raw_html)
        if price:
            return price, f"fpt_{method}"
        return first_visible_price_after(visible, ("Giá tại",)), "fpt_visible_price"
    if retailer == "CPS":
        price, method = extract_cps_price(raw_html)
        if price:
            return price, method
        return first_visible_price_after(visible, ("Hoặc", "Phiên bản", "Màu sắc")), "cps_visible_fallback"
    if retailer == "Shopdunk":
        return extract_shopdunk_price(raw_html)
    return None, ""


def final_price_result(
    target: Target,
    price: int,
    visible: str,
    fetched_at: str,
    source_method: str,
    purchase_action: tuple[str, str, str] | None = None,
) -> Result:
    discount, promo_note = eligible_direct_discount(target.retailer, target.model, visible)
    final_price = max(price - discount, 0)
    result = Result(
        target.model,
        target.retailer,
        target.url,
        final_price,
        "OK",
        fetched_at=fetched_at,
        base_value=price,
        discount=discount,
        promo_note=promo_note,
        source_method=source_method,
        html_price=final_price,
    )
    if purchase_action:
        result.purchase_action, result.purchase_action_method, result.purchase_action_text = purchase_action
    return result


def extract_price(raw_html: str, target: Target) -> Result:
    fetched_at = dt.datetime.now().isoformat(timespec="seconds")
    visible = strip_tags(raw_html)
    purchase_action = extract_purchase_action(raw_html)

    if normalize_name(target.retailer) == "Viettel":
        display_price, source_method = extract_viettel_product_offer(raw_html, target.url)
    else:
        display_price, source_method = extract_site_display_price(raw_html, visible, target.retailer, target.model)
    oos = detect_product_oos(raw_html, visible, target.retailer, display_price, target.url)
    if oos:
        return oos_result(target, fetched_at, *oos)
    if display_price:
        return final_price_result(target, display_price, visible, fetched_at, source_method, purchase_action)

    structured_price, structured_method = extract_structured_price(raw_html)
    oos = detect_product_oos(raw_html, visible, target.retailer, structured_price, target.url)
    if oos:
        return oos_result(target, fetched_at, *oos)
    if structured_price:
        return final_price_result(target, structured_price, visible, fetched_at, structured_method, purchase_action)

    candidates = []
    for match in re.finditer(r"([0-9]{1,3}(?:[.,][0-9]{3}){2,})\s*(?:đ|₫|vnd)", visible, flags=re.I):
        price = parse_price_number(match.group(1))
        if price and 500_000 <= price <= 100_000_000:
            candidates.append(price)
    if candidates:
        oos = detect_product_oos(raw_html, visible, target.retailer, candidates[0], target.url)
        if oos:
            return oos_result(target, fetched_at, *oos)
        return final_price_result(target, candidates[0], visible, fetched_at, "visible_first_price", purchase_action)
    oos = detect_product_oos(raw_html, visible, target.retailer, None, target.url)
    if oos:
        return oos_result(target, fetched_at, *oos)
    return Result(target.model, target.retailer, target.url, None, "NO_PRICE", "No price pattern matched", fetched_at=fetched_at)


def fetch_target(target: Target, delay: float, overrides: dict[str, tuple[int, str]] | None = None) -> Result:
    fetched_at = dt.datetime.now().isoformat(timespec="seconds")
    try:
        result = extract_price(request_html(target.url), target)
    except HTTPError as exc:
        result = Result(target.model, target.retailer, target.url, None, "HTTP_ERROR", f"{exc.code} {exc.reason}", fetched_at=fetched_at)
    except URLError as exc:
        result = Result(target.model, target.retailer, target.url, None, "URL_ERROR", str(exc.reason), fetched_at=fetched_at)
    except Exception as exc:  # noqa: BLE001 - keep batch running and report URL failures.
        result = Result(target.model, target.retailer, target.url, None, "ERROR", str(exc), fetched_at=fetched_at)
    if result.status not in {"HTTP_ERROR", "URL_ERROR", "ERROR"} and normalize_name(target.retailer) == "Shopdunk":
        result = apply_shopdunk_stock_probe(result, target)
    if result.status != "OOS" and result.status not in {"HTTP_ERROR", "URL_ERROR", "ERROR"} and overrides and target.url in overrides:
        value, note = overrides[target.url]
        result = Result(
            target.model,
            target.retailer,
            target.url,
            value,
            "OK",
            error="Manual override",
            fetched_at=fetched_at,
            base_value=value,
            promo_note=note,
            source_method="override",
            html_price=value,
            confidence="OVERRIDE",
            risk_flags="OVERRIDE_USED",
            decision_note=note,
        )
    if delay:
        time.sleep(delay)
    return result


def result_key(result: Result) -> str:
    return f"{result.model}||{normalize_name(result.retailer)}"


def numeric_price(value: int | str | None) -> int | None:
    return value if isinstance(value, int) else None


def median(values: list[int]) -> int | None:
    if not values:
        return None
    ordered = sorted(values)
    middle = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[middle]
    return (ordered[middle - 1] + ordered[middle]) // 2


def suspicious_low_price(model: str, value: int) -> bool:
    lower = model.lower()
    thresholds = (
        ("macbook", 10_000_000),
        ("ipad pro", 10_000_000),
        ("iphone", 5_000_000),
        ("aw ", 3_000_000),
        ("apple watch", 3_000_000),
        ("airpods pro", 2_000_000),
        ("airpods", 1_000_000),
    )
    return any(token in lower and value < threshold for token, threshold in thresholds)


def is_review_needed(result: Result) -> bool:
    if result.status == "OOS" and result.confidence in {"VERIFIED_OOS", "OOS_CONFIRMED"}:
        return False
    flags = set(filter(None, result.risk_flags.split("|")))
    return result.status not in {"OK", "OOS"} or bool(flags & REVIEW_STATUS_FLAGS)


def is_render_candidate(result: Result) -> bool:
    if result.source_method == "override":
        return False
    retailer = normalize_name(result.retailer)
    render_recovery_retailer = retailer in {"CPS", "MW", "Shopdunk", "FPT", "Viettel"}
    if result.source_method == "missing_direct_link":
        return False
    weak_render_retailer = retailer in {"CPS", "MW", "Shopdunk"}
    if result.status == "OOS":
        return render_recovery_retailer
    if result.status != "OK":
        return render_recovery_retailer
    if not isinstance(result.value, int):
        return False
    flags = set(filter(None, result.risk_flags.split("|")))
    if retailer == "Viettel" and result.source_method == "viettel_exact_jsonld_offer":
        return False
    if retailer == "Viettel" and result.source_method == "structured_offer_price":
        return True
    if retailer == "FPT" and result.source_method in {"fpt_structured_offer_price", "structured_offer_price"}:
        return True
    weak_source = result.source_method in {
        "structured_offer_price",
        "structured_product_price",
        "structured_itemprop_price",
        "structured_meta_price",
        "structured_json_price",
        "jsonld_min_price",
        "visible_first_price",
        "cps_schema_offer",
        "cps_meta_price",
        "cps_json_price",
        "fpt_visible_price",
    }
    return bool(flags & {"WARN_WEEK_CHANGE", "WARN_OUTLIER", "WARN_SUSPICIOUS_LOW"}) or (weak_render_retailer and weak_source)


def is_trusted_render_method(method: str) -> bool:
    return any(method.startswith(prefix) for prefix in TRUSTED_RENDER_METHOD_PREFIXES)


def is_trusted_oos_method(method: str) -> bool:
    return method in {
        "structured_offer_availability",
        "product_status_element",
        "mw_price_context",
        "shopdunk_stock_endpoint",
    } or method.startswith(
        ("html_action:", "rendered_dom:", "rendered_price_context")
    )


def workbook_context(workbook_path: Path, source_sheet: str) -> tuple[dict[tuple[str, str], int | str | None], dict[str, list[int]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    source = wb[source_sheet]
    retailer_to_col: dict[str, int] = {}
    for col in range(2, source.max_column + 1):
        header = source.cell(1, col).value
        if header:
            retailer_to_col[normalize_name(str(header))] = col

    previous: dict[tuple[str, str], int | str | None] = {}
    model_channel_values: dict[str, list[int]] = {}
    for row in range(2, source.max_row + 1):
        model = str(source.cell(row, 1).value or "").strip()
        if not model:
            continue
        for retailer, col in retailer_to_col.items():
            value = source.cell(row, col).value
            previous[(model, retailer)] = value
            if isinstance(value, int):
                model_channel_values.setdefault(model, []).append(value)
    return previous, model_channel_values


def add_initial_risk_flags_from_previous(
    results: list[Result],
    previous: dict[tuple[str, str], int | str | None],
) -> None:
    """Apply the shared crawler risk rules against a supplied prior-price map."""
    current_by_model: dict[str, list[int]] = {}
    for result in results:
        if isinstance(result.value, int):
            current_by_model.setdefault(result.model, []).append(result.value)

    for result in results:
        retailer = normalize_name(result.retailer)
        result.previous_week_price = previous.get((result.model, retailer))
        flags = set(filter(None, result.risk_flags.split("|")))
        notes = []

        if result.status not in {"OK", "OOS"}:
            flags.add("WARN_FETCH_STATUS")
            notes.append(result.error or result.status)

        if result.status == "OOS" and result.previous_week_price != "OOS":
            flags.add("WARN_WEEK_CHANGE")

        if isinstance(result.value, int):
            if result.previous_week_price != result.value:
                flags.add("WARN_WEEK_CHANGE")
            if suspicious_low_price(result.model, result.value):
                flags.add("WARN_SUSPICIOUS_LOW")
                notes.append("price below product-family floor")
            peers = [value for value in current_by_model.get(result.model, []) if value != result.value]
            peer_median = median(peers)
            if peer_median and (result.value < peer_median * 0.7 or result.value > peer_median * 1.3):
                flags.add("WARN_OUTLIER")
                notes.append(f"outlier vs peer median {format_price(peer_median)}")

        weak_source = result.source_method in {
            "structured_offer_price",
            "structured_product_price",
            "structured_itemprop_price",
            "structured_meta_price",
            "structured_json_price",
            "jsonld_min_price",
            "visible_first_price",
            "cps_schema_offer",
            "cps_meta_price",
            "cps_json_price",
            "fpt_visible_price",
        }
        if weak_source:
            flags.add("WARN_WEAK_SOURCE")
        if (
            result.status == "OK"
            and result.source_method != "override"
            and retailer in {"FPT", "Viettel", "MW", "CPS", "Shopdunk"}
            and result.purchase_action != "IN_STOCK"
        ):
            flags.add("WARN_STOCK_ACTION_UNVERIFIED")

        result.risk_flags = "|".join(sorted(flags))
        if result.confidence not in {"OVERRIDE", "OOS_CONFIRMED"}:
            result.confidence = "REVIEW" if is_review_needed(result) else "HTML_OK"
        if notes and not result.decision_note:
            result.decision_note = "; ".join(notes)


def add_initial_risk_flags(results: list[Result], workbook_path: Path, source_sheet: str) -> None:
    previous, _ = workbook_context(workbook_path, source_sheet)
    add_initial_risk_flags_from_previous(results, previous)


def run_render_backcheck(results: list[Result], report_dir: Path, enabled: bool = True) -> None:
    if not enabled:
        return
    candidates = [result for result in results if is_render_candidate(result)]
    if not candidates:
        return

    runtime = resolve_browser_runtime()
    helper = Path(__file__).with_name("render_price_check.cjs")
    if not (runtime.node_bin and runtime.node_bin.exists() and helper.exists()):
        for result in candidates:
            flags = set(filter(None, result.risk_flags.split("|")))
            flags.add("WARN_RENDER_UNAVAILABLE")
            result.risk_flags = "|".join(sorted(flags))
            result.confidence = "REVIEW"
            result.decision_note = "browser render back-check unavailable"
        return

    report_dir.mkdir(parents=True, exist_ok=True)
    env = dict(os.environ)
    if runtime.node_path:
        env["NODE_PATH"] = str(runtime.node_path)
    rendered_items: list[dict] = []
    render_errors: dict[str, str] = {}
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    for batch_index, start in enumerate(range(0, len(candidates), RENDER_BATCH_SIZE), start=1):
        batch = candidates[start : start + RENDER_BATCH_SIZE]
        input_path = report_dir / f"render_tasks_{stamp}_b{batch_index}.json"
        tasks = [
            {"key": result_key(result), "url": result.url, "retailer": normalize_name(result.retailer), "model": result.model}
            for result in batch
        ]
        input_path.write_text(json.dumps(tasks, ensure_ascii=False), encoding="utf-8")
        try:
            command = [str(runtime.node_bin), str(helper), str(input_path)]
            if runtime.chromium_bin:
                command.append(str(runtime.chromium_bin))
            completed = subprocess.run(
                command,
                cwd=str(Path.cwd()),
                env=env,
                text=True,
                capture_output=True,
                # Each page can spend up to 45s in page.goto; allow the full
                # batch to finish so a slow site does not mark later URLs as
                # review-needed just because the subprocess was cut short.
                timeout=max(120, len(tasks) * 60),
                check=False,
            )
            if completed.returncode == 0:
                rendered_items.extend(json.loads(completed.stdout or "[]"))
                continue
            render_error = completed.stderr.strip() or "browser render returned no price"
        except Exception as exc:  # noqa: BLE001 - report and keep workbook update moving.
            render_error = str(exc)
        finally:
            input_path.unlink(missing_ok=True)
        for result in batch:
            render_errors[result_key(result)] = render_error

    rendered_by_key = {item.get("key"): item for item in rendered_items}
    for result in candidates:
        flags = set(filter(None, result.risk_flags.split("|")))
        item = rendered_by_key.get(result_key(result))
        if not item or not item.get("ok"):
            if result.status == "OOS" and is_trusted_oos_method(result.availability_method):
                flags.discard("WARN_RENDER_UNAVAILABLE")
                result.confidence = "OOS_CONFIRMED"
                result.decision_note = f"trusted stock signal reported OOS via {result.availability_method}"
            elif (
                normalize_name(result.retailer) == "Shopdunk"
                and result.source_method == "shopdunk_price_value"
                and isinstance(result.value, int)
                and result.purchase_action == "IN_STOCK"
            ):
                flags.discard("WARN_RENDER_UNAVAILABLE")
                result.confidence = "HTML_OK"
                result.decision_note = "trusted Shopdunk product price; browser back-check unavailable"
            else:
                flags.add("WARN_RENDER_UNAVAILABLE")
                result.confidence = "REVIEW"
                batch_error = render_errors.get(result_key(result), "")
                result.decision_note = result.decision_note or item.get("error", "") if item else batch_error or "browser render returned no price"
        else:
            rendered_price = parse_price_number(str(item.get("price") or ""))
            render_method = str(item.get("method") or "")
            render_action_status = str(item.get("actionStatus") or "")
            render_action_method = str(item.get("actionMethod") or "")
            render_action_text = str(item.get("actionText") or "")
            render_availability = str(item.get("availability") or "")
            render_availability_method = str(item.get("availabilityMethod") or "")
            render_availability_text = str(item.get("availabilityText") or "")
            result.rendered_price = rendered_price
            if render_action_status:
                result.purchase_action = render_action_status
                result.purchase_action_method = render_action_method
                result.purchase_action_text = compact_evidence(render_action_text)
            if render_availability == "OOS":
                result.value = "OOS"
                result.base_value = "OOS"
                result.status = "OOS"
                result.error = ""
                result.availability = "OOS"
                result.availability_method = render_availability_method
                result.availability_text = compact_evidence(render_availability_text)
                result.source_method = result.source_method or render_availability_method
                flags.discard("WARN_FETCH_STATUS")
                flags.add("WARN_RENDER_STOCK_CONFIRMED")
                if result.previous_week_price != "OOS":
                    flags.add("WARN_WEEK_CHANGE")
                result.confidence = "VERIFIED_OOS"
                result.decision_note = f"render confirmed OOS via {render_availability_method}"
            elif result.status == "OOS":
                # A retailer's structured availability can describe another or
                # stale variant. A rendered price in the selected SKU element is
                # stronger evidence that this exact link is purchasable.
                retailer = normalize_name(result.retailer)
                trusted_selected_price = (
                    (retailer == "MW" and render_method == "rendered_dom:.box-price-present")
                    or (retailer == "Viettel" and render_method == "rendered_dom:.version-product.active .txt-price")
                )
                trusted_cps_price = retailer == "CPS" and render_method in {
                    "rendered_dom:.sale-price",
                    "rendered_dom:.box-product-price",
                }
                trusted_purchase_action = render_action_status == "IN_STOCK" and bool(rendered_price)
                if (
                    rendered_price
                    and (
                        trusted_purchase_action
                        or
                        (result.availability_method == "structured_offer_availability" and trusted_selected_price)
                        or (result.availability_method == "cps_price_context" and trusted_cps_price)
                    )
                ):
                    result.value = rendered_price
                    result.base_value = rendered_price
                    result.status = "OK"
                    result.error = ""
                    result.availability = ""
                    result.availability_method = ""
                    result.availability_text = ""
                    result.source_method = "trusted_render_price"
                    flags.discard("WARN_FETCH_STATUS")
                    flags.discard("WARN_HTML_RENDER_MISMATCH")
                    if result.previous_week_price != rendered_price:
                        flags.add("WARN_WEEK_CHANGE")
                    result.confidence = "VERIFIED_RENDER"
                    result.decision_note = (
                        f"{retailer} rendered selected-SKU price {format_price(rendered_price)} via {render_method}; "
                        "discarded stale structured availability"
                    )
                elif is_trusted_oos_method(result.availability_method):
                    flags.discard("WARN_HTML_RENDER_MISMATCH")
                    result.confidence = "OOS_CONFIRMED"
                    result.decision_note = (
                        f"trusted stock signal reported OOS via {result.availability_method}; "
                        "render retained price without stock text"
                    )
                else:
                    flags.add("WARN_HTML_RENDER_MISMATCH")
                    result.confidence = "REVIEW"
                    result.decision_note = (
                        f"html reported OOS via {result.availability_method}; render did not confirm stock status"
                    )
            elif rendered_price and result.status != "OK":
                result.value = rendered_price
                result.base_value = rendered_price
                result.status = "OK"
                result.error = ""
                result.source_method = "rendered_recovery"
                flags.discard("WARN_FETCH_STATUS")
                flags.add("WARN_RENDER_RECOVERED")
                if result.previous_week_price != rendered_price:
                    flags.add("WARN_WEEK_CHANGE")
                result.confidence = "VERIFIED_RENDER"
                result.decision_note = f"html failed; render recovered {format_price(rendered_price)} via {render_method}"
            elif rendered_price and isinstance(result.value, int):
                if render_action_status == "IN_STOCK" or result.purchase_action == "IN_STOCK":
                    flags.discard("WARN_STOCK_ACTION_UNVERIFIED")
                else:
                    flags.add("WARN_STOCK_ACTION_UNVERIFIED")
                acceptable = {result.value}
                if isinstance(result.base_value, int):
                    acceptable.add(result.base_value)
                if rendered_price in acceptable:
                    if is_trusted_render_method(render_method):
                        flags.discard("WARN_OUTLIER")
                        flags.discard("WARN_WEAK_SOURCE")
                        flags.discard("WARN_SUSPICIOUS_LOW")
                    if not (flags & REVIEW_STATUS_FLAGS):
                        result.confidence = "VERIFIED_RENDER"
                    result.decision_note = result.decision_note or f"render matched via {render_method}"
                elif not is_trusted_render_method(render_method):
                    result.decision_note = result.decision_note or (
                        f"ignored untrusted render {format_price(rendered_price)} via {render_method}"
                    )
                else:
                    html_price = result.value
                    result.value = rendered_price
                    result.base_value = rendered_price
                    result.status = "OK"
                    result.error = ""
                    result.source_method = "trusted_render_price"
                    flags.discard("WARN_OUTLIER")
                    flags.discard("WARN_WEAK_SOURCE")
                    flags.discard("WARN_SUSPICIOUS_LOW")
                    flags.discard("WARN_HTML_RENDER_MISMATCH")
                    if result.previous_week_price != rendered_price:
                        flags.add("WARN_WEEK_CHANGE")
                    result.confidence = "VERIFIED_RENDER"
                    result.decision_note = (
                        f"trusted render replaced html {format_price(html_price)} "
                        f"with {format_price(rendered_price)} via {render_method}"
                    )
        result.risk_flags = "|".join(sorted(flags))


def write_review_report(results: list[Result], report_dir: Path, label: str) -> Path | None:
    review_rows = [result for result in results if is_review_needed(result)]
    if not review_rows:
        return None
    report_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    path = report_dir / f"review_needed_{label}_{stamp}.csv"
    write_results_csv(review_rows, path)
    return path


def write_results_csv(results: list[Result], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "model",
                "retailer",
                "value",
                "final_price",
                "html_price",
                "rendered_price",
                "previous_week_price",
                "base_value",
                "discount",
                "promo_note",
                "source_method",
                "availability",
                "availability_method",
                "availability_text",
                "purchase_action",
                "purchase_action_method",
                "purchase_action_text",
                "risk_flags",
                "confidence",
                "decision_note",
                "status",
                "error",
                "fetched_at",
                "url",
            ],
        )
        writer.writeheader()
        for result in results:
            row = result.__dict__.copy()
            row["final_price"] = result.value
            writer.writerow(row)


def write_report(results: list[Result], report_dir: Path, label: str) -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    path = report_dir / f"price_results_{label}_{stamp}.csv"
    write_results_csv(results, path)
    return path


def update_formula_refs(formula: str, old_sheet: str | None, new_ref_sheet: str) -> str:
    if not old_sheet:
        return formula
    return formula.replace(f"{quote_sheetname(old_sheet)}!", f"{quote_sheetname(new_ref_sheet)}!").replace(
        f"{old_sheet}!", f"{quote_sheetname(new_ref_sheet)}!"
    )


def detect_previous_sheet_ref(ws) -> str | None:
    counts: dict[str, int] = {}
    sheet_names = [name for name in ws.parent.sheetnames if name != ws.title]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                for name in sheet_names:
                    if f"{name}!" in cell.value or f"{quote_sheetname(name)}!" in cell.value:
                        counts[name] = counts.get(name, 0) + 1
    if not counts:
        return None
    return max(counts, key=counts.get)


def count_formula_sheet_refs(ws, sheet_name: str) -> int:
    quoted_ref = f"{quote_sheetname(sheet_name)}!"
    plain_ref = f"{sheet_name}!"
    count = 0
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or not cell.value.startswith("="):
                continue
            count += cell.value.count(quoted_ref)
            if quoted_ref != plain_ref:
                count += cell.value.count(plain_ref)
    return count


def validate_inherited_formula_refs(target, old_sheet: str | None, source_sheet: str, expected_refs: int) -> None:
    if not old_sheet or expected_refs == 0:
        return
    stale_refs = count_formula_sheet_refs(target, old_sheet)
    source_refs = count_formula_sheet_refs(target, source_sheet)
    if stale_refs or source_refs < expected_refs:
        raise ValueError(
            "Formula inheritance validation failed: "
            f"expected at least {expected_refs} references to {source_sheet}, "
            f"found {source_refs}; stale references to {old_sheet}: {stale_refs}."
        )


def fill_key(fill: PatternFill) -> tuple[str | None, str, str]:
    return (fill.fill_type, fill.fgColor.type, fill.fgColor.rgb or fill.fgColor.indexed or fill.fgColor.theme or "")


def is_changed_fill(fill: PatternFill) -> bool:
    return fill.fill_type == "solid" and (fill.fgColor.rgb or "").upper() in {"FFFFFF00", "FFFF00", "00FFFF00"}


def prepare_target_sheet(wb, source_sheet: str, target_sheet: str):
    source = wb[source_sheet]
    previous_ref = detect_previous_sheet_ref(source)
    expected_refs = count_formula_sheet_refs(source, previous_ref) if previous_ref else 0
    if target_sheet in wb.sheetnames:
        target = wb[target_sheet]
        for row in source.iter_rows():
            for source_cell in row:
                target_cell = target[source_cell.coordinate]
                if target_cell.value is None and isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                    target_cell.value = update_formula_refs(source_cell.value, previous_ref, source_sheet)
        validate_inherited_formula_refs(target, previous_ref, source_sheet, expected_refs)
        return target, source

    target = wb.copy_worksheet(source)
    target.title = target_sheet
    for row in target.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = update_formula_refs(cell.value, previous_ref, source_sheet)
    validate_inherited_formula_refs(target, previous_ref, source_sheet, expected_refs)
    return target, source


def reset_price_area_comments_and_highlights(
    target,
    retailer_to_col: dict[str, int],
    retailers_to_reset: set[str],
    rows_to_reset: set[int],
) -> None:
    neutral_fills: dict[tuple[str | None, str, str], PatternFill] = {}
    fill_counts: dict[tuple[str | None, str, str], int] = {}
    for retailer, col in retailer_to_col.items():
        if retailer not in retailers_to_reset:
            continue
        for row in rows_to_reset:
            if not target.cell(row, 1).value or target.cell(row, col).value is None:
                continue
            fill = target.cell(row, col).fill
            if is_changed_fill(fill):
                continue
            key = fill_key(fill)
            neutral_fills[key] = fill
            fill_counts[key] = fill_counts.get(key, 0) + 1

    neutral_fill = neutral_fills[max(fill_counts, key=fill_counts.get)] if fill_counts else None
    for retailer, col in retailer_to_col.items():
        if retailer not in retailers_to_reset:
            continue
        for row in rows_to_reset:
            target.cell(row, col).comment = None
            if neutral_fill is not None and target.cell(row, 1).value:
                target.cell(row, col).fill = copy.copy(neutral_fill)


def update_workbook(
    workbook_path: Path,
    source_sheet: str,
    target_sheet: str,
    results: list[Result],
    backup_dir: Path,
) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"{workbook_path.stem}_backup_{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    shutil.copy2(workbook_path, backup_path)

    wb = load_workbook(workbook_path)
    target, source = prepare_target_sheet(wb, source_sheet, target_sheet)

    retailer_to_col: dict[str, int] = {}
    for col in range(2, target.max_column + 1):
        header = target.cell(1, col).value
        if header:
            retailer_to_col[normalize_name(str(header))] = col
    model_to_row = {str(target.cell(row, 1).value).strip(): row for row in range(2, target.max_row + 1) if target.cell(row, 1).value}

    retailers_to_update = {normalize_name(result.retailer) for result in results if result.status in {"OK", "OOS"}}
    rows_to_update = {model_to_row[result.model] for result in results if result.status in {"OK", "OOS"} and result.model in model_to_row}
    reset_price_area_comments_and_highlights(target, retailer_to_col, retailers_to_update, rows_to_update)

    updated = 0
    for result in results:
        if result.status not in {"OK", "OOS"}:
            continue
        row = model_to_row.get(result.model)
        col = retailer_to_col.get(normalize_name(result.retailer))
        if not row or not col:
            continue
        cell = target.cell(row, col)
        previous_value = source.cell(row, col).value
        cell.value = result.value if result.status == "OK" else "OOS"
        price_changed = previous_value != cell.value
        needs_review = is_review_needed(result)
        if needs_review:
            cell.fill = PRICE_REVIEW_FILL
        elif price_changed:
            cell.fill = PRICE_CHANGED_FILL
        retailer = normalize_name(result.retailer)
        should_comment = price_changed or needs_review or result.source_method == "override" or (result.discount and retailer != "MW")
        if should_comment:
            comment_lines = [f"was {format_price(previous_value)}"]
            if result.discount and retailer != "MW":
                comment_lines.append(f"base {format_price(result.base_value)}")
                comment_lines.append(f"discount {format_price(result.discount)}")
                if result.promo_note:
                    comment_lines.append(result.promo_note)
            if needs_review or result.source_method == "override":
                if result.html_price:
                    comment_lines.append(f"html {format_price(result.html_price)}")
                if result.rendered_price:
                    comment_lines.append(f"rendered {format_price(result.rendered_price)}")
                if result.risk_flags:
                    comment_lines.append(f"risk {result.risk_flags}")
                if result.confidence:
                    comment_lines.append(f"confidence {result.confidence}")
                if result.decision_note:
                    comment_lines.append(result.decision_note)
            cell.comment = Comment("\n".join(comment_lines), "Price Check Tool")
        updated += 1

    wb.save(workbook_path)
    print(f"Updated {workbook_path} sheet {target_sheet}; wrote {updated} cells.")
    print(f"Backup: {backup_path}")
    return backup_path


def run(args: argparse.Namespace) -> None:
    workbook_path = Path(args.workbook)
    link_workbook = Path(args.link_workbook)
    retailer_filter = {normalize_name(r) for r in args.retailers.split(",")} if args.retailers else None
    targets = read_link_targets(link_workbook, retailer_filter)
    all_link_targets = read_link_targets(link_workbook)
    overrides = read_price_overrides(Path(args.override_file))
    if args.limit:
        targets = targets[: args.limit]
    if not targets:
        print("No direct product links found in Link Product workbook.", file=sys.stderr)
        return

    source_sheet = args.source_sheet or latest_week_sheet(workbook_sheets(workbook_path))
    target_sheet = args.target_sheet or next_week_name(source_sheet)
    print(f"Source sheet: {source_sheet}")
    print(f"Target sheet: {target_sheet}")
    print(f"Targets: {len(targets)}")

    results: list[Result] = []
    for index, target in enumerate(targets, start=1):
        print(f"[{index}/{len(targets)}] {target.retailer}: {target.model}")
        results.append(fetch_target(target, args.delay, overrides))

    results.extend(missing_link_results(workbook_path, source_sheet, link_workbook, all_link_targets))
    add_initial_risk_flags(results, workbook_path, source_sheet)
    run_render_backcheck(results, Path(args.report_dir), enabled=not args.no_back_check)
    report_path = write_report(results, Path(args.report_dir), target_sheet)
    print(f"Report: {report_path}")
    review_path = write_review_report(results, Path(args.report_dir), target_sheet)
    if review_path:
        print(f"Review needed: {review_path}")
    if args.dry_run:
        return
    if args.require_clean and any(is_review_needed(result) for result in results):
        raise SystemExit("Refusing workbook update because review-needed rows remain.")
    update_workbook(workbook_path, source_sheet, target_sheet, results, Path(args.backup_dir))


def main() -> None:
    parser = argparse.ArgumentParser(description="Update Price Check workbook from direct product links.")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--link-workbook", default=str(DEFAULT_LINK_WORKBOOK))
    parser.add_argument("--override-file", default=str(DEFAULT_OVERRIDE_FILE))
    parser.add_argument("--source-sheet")
    parser.add_argument("--target-sheet")
    parser.add_argument("--retailers", help="Comma-separated retailer filter, e.g. Viettel,MW,FPT,CPS,Shopdunk")
    parser.add_argument("--report-dir", default=str(DEFAULT_REPORT_DIR))
    parser.add_argument("--backup-dir", default=str(DEFAULT_BACKUP_DIR))
    parser.add_argument("--delay", type=float, default=0.5)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-back-check", action="store_true", help="Skip browser-render back-check for risky rows.")
    parser.add_argument("--require-clean", action="store_true", help="Refuse workbook updates while review-needed rows remain.")
    args = parser.parse_args()
    run(args)


if __name__ == "__main__":
    main()
