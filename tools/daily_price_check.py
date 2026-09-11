#!/usr/bin/env python3
"""Run daily price checks without modifying the weekly Price Check workbook."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import csv
import datetime as dt
import json
import os
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


ROOT = Path(__file__).resolve().parents[1]
LOCAL_TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")
DEFAULT_DAILY_WORKBOOK = ROOT / "Price Check Daily.xlsx"
DEFAULT_WEEKLY_WORKBOOK = ROOT / "Price Check.xlsx"
DEFAULT_LINK_WORKBOOK = ROOT / "Link Product.xlsx"
DEFAULT_OVERRIDE_FILE = ROOT / "Price Overrides.csv"
DEFAULT_DAILY_DIR = ROOT / "outputs" / "daily"
PRICE_CHECK_PATH = ROOT / "tools" / "price_check_tool.py"
PRICE_SPEC = importlib.util.spec_from_file_location("price_check_tool", PRICE_CHECK_PATH)
if PRICE_SPEC is None or PRICE_SPEC.loader is None:  # pragma: no cover - installation failure
    raise ImportError(f"Cannot load {PRICE_CHECK_PATH}")
price_check = importlib.util.module_from_spec(PRICE_SPEC)
sys.modules.setdefault(PRICE_SPEC.name, price_check)
PRICE_SPEC.loader.exec_module(price_check)

LATEST_HEADERS = [
    "Key", "Model", "Retailer", "Current Value", "Stock Status", "Previous Value",
    "Delta VND", "First Seen", "Last Changed", "Last Checked", "Confidence",
    "Risk Flags", "Review Status", "Source Method", "URL", "Run ID",
]
CHANGE_HEADERS = [
    "Event ID", "Detected At", "Model", "Retailer", "Change Type", "Old Value",
    "New Value", "Delta VND", "Delta Percent", "Confidence", "Risk Flags",
    "Review Status", "Source Method", "URL", "Run ID",
]
RUN_LOG_HEADERS = [
    "Run ID", "Parent Run ID", "Run Type", "Started At", "Completed At",
    "Total Targets", "OK", "OOS", "Errors", "Changes", "Review Needed",
    "Status", "Snapshot Path", "Review Path", "Backup Path",
]
REVIEW_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")

LATEST_FIELD_MAP = {
    "Key": "key", "Model": "model", "Retailer": "retailer", "Current Value": "current_value",
    "Stock Status": "stock_status", "Previous Value": "previous_value", "Delta VND": "delta_vnd",
    "First Seen": "first_seen", "Last Changed": "last_changed", "Last Checked": "last_checked",
    "Confidence": "confidence", "Risk Flags": "risk_flags", "Review Status": "review_status",
    "Source Method": "source_method", "URL": "url", "Run ID": "run_id",
}
CHANGE_FIELD_MAP = {
    "Event ID": "event_id", "Detected At": "detected_at", "Model": "model", "Retailer": "retailer",
    "Change Type": "change_type", "Old Value": "old_value", "New Value": "new_value",
    "Delta VND": "delta_vnd", "Delta Percent": "delta_percent", "Confidence": "confidence",
    "Risk Flags": "risk_flags", "Review Status": "review_status", "Source Method": "source_method",
    "URL": "url", "Run ID": "run_id",
}


@dataclass
class Reconciliation:
    latest: dict[str, dict]
    changes: list[dict]
    failures: list[price_check.Result]


@dataclass(frozen=True)
class RunMetadata:
    run_id: str
    parent_run_id: str
    run_type: str
    started_at: str
    completed_at: str
    snapshot_path: str
    review_path: str = ""


@dataclass(frozen=True)
class WorkbookSummary:
    run_id: str
    total_targets: int
    ok_count: int
    oos_count: int
    error_count: int
    change_count: int
    review_count: int
    backup_path: str
    already_recorded: bool = False
    changes: tuple[dict, ...] = ()


def observation_key(model: str, retailer: str) -> str:
    return f"{model.strip()}||{price_check.normalize_name(retailer)}"


def _stock_status(result: price_check.Result) -> str:
    return "OOS" if result.status == "OOS" or str(result.value).upper() == "OOS" else "In stock"


def _is_success(result: price_check.Result) -> bool:
    return result.status in {"OK", "OOS"}


def _is_risky(result: price_check.Result) -> bool:
    return price_check.is_review_needed(result)


def _event_id(run_id: str, key: str, old_value, new_value, old_stock: str, new_stock: str) -> str:
    raw = "|".join(map(str, (run_id, key, old_value, new_value, old_stock, new_stock)))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:24]


def reconcile_results(
    previous: dict[str, dict],
    results: list[price_check.Result],
    run_id: str,
    observed_at: str,
    existing_event_ids: set[str] | None = None,
) -> Reconciliation:
    """Apply successful observations and return immutable change events."""
    latest = {key: value.copy() for key, value in previous.items()}
    changes: list[dict] = []
    failures: list[price_check.Result] = []
    seen_event_ids = existing_event_ids or set()

    for result in results:
        key = observation_key(result.model, result.retailer)
        if not _is_success(result):
            failures.append(result)
            continue

        old = latest.get(key)
        stock = _stock_status(result)
        current_value = "OOS" if stock == "OOS" else result.value
        record = {
            "key": key,
            "model": result.model,
            "retailer": price_check.normalize_name(result.retailer),
            "current_value": current_value,
            "stock_status": stock,
            "previous_value": old.get("current_value") if old else None,
            "delta_vnd": None,
            "first_seen": old.get("first_seen", observed_at) if old else observed_at,
            "last_changed": old.get("last_changed", observed_at) if old else observed_at,
            "last_checked": observed_at,
            "confidence": result.confidence,
            "risk_flags": result.risk_flags,
            "review_status": "Pending" if _is_risky(result) else "",
            "source_method": result.source_method,
            "url": result.url,
            "run_id": run_id,
        }

        if old:
            old_value = old.get("current_value")
            old_stock = old.get("stock_status", "In stock")
            changed = old_value != current_value or old_stock != stock
            if isinstance(old_value, (int, float)) and isinstance(current_value, (int, float)):
                record["delta_vnd"] = int(current_value - old_value)
            if changed:
                record["last_changed"] = observed_at
                if old_stock != stock:
                    change_type = f"Stock: {old_stock} to {stock}"
                elif current_value < old_value:
                    change_type = "Price decrease"
                else:
                    change_type = "Price increase"
                event_id = _event_id(run_id, key, old_value, current_value, old_stock, stock)
                if event_id not in seen_event_ids:
                    delta = record["delta_vnd"]
                    delta_percent = delta / old_value if delta is not None and old_value else None
                    changes.append(
                        {
                            "event_id": event_id,
                            "detected_at": observed_at,
                            "model": result.model,
                            "retailer": price_check.normalize_name(result.retailer),
                            "change_type": change_type,
                            "old_value": old_value,
                            "new_value": current_value,
                            "delta_vnd": delta,
                            "delta_percent": delta_percent,
                            "confidence": result.confidence,
                            "risk_flags": result.risk_flags,
                            "review_status": "Pending" if _is_risky(result) else "",
                            "source_method": result.source_method,
                            "url": result.url,
                            "run_id": run_id,
                        }
                    )
        latest[key] = record

    return Reconciliation(latest=latest, changes=changes, failures=failures)


def _iso_cell(value):
    if not isinstance(value, str) or "T" not in value:
        return value
    try:
        parsed = dt.datetime.fromisoformat(value)
        return parsed.replace(tzinfo=None) if parsed.tzinfo is not None else parsed
    except ValueError:
        return value


def _iso_text(value):
    return value.isoformat() if isinstance(value, (dt.datetime, dt.date)) else value


def _read_latest(sheet) -> dict[str, dict]:
    headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
    latest = {}
    for row in range(2, sheet.max_row + 1):
        key = sheet.cell(row, headers.get("Key", 1)).value
        if not key:
            continue
        record = {}
        for header, field in LATEST_FIELD_MAP.items():
            col = headers.get(header)
            record[field] = _iso_text(sheet.cell(row, col).value) if col else None
        latest[str(key)] = record
    return latest


def _existing_values(sheet, header_name: str) -> set[str]:
    headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
    col = headers.get(header_name)
    if not col:
        return set()
    return {str(sheet.cell(row, col).value) for row in range(2, sheet.max_row + 1) if sheet.cell(row, col).value}


def _compact_data_rows(sheet, key_column: int = 1) -> None:
    """Remove styled-but-empty gaps while preserving rows with a real key."""
    data = [
        [sheet.cell(row, col).value for col in range(1, sheet.max_column + 1)]
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, key_column).value not in {None, ""}
    ]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in data:
        sheet.append(row)


def _paint_pending_rows(sheet, review_header: str = "Review Status") -> None:
    headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
    review_col = headers.get(review_header)
    if not review_col:
        return
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, review_col).value == "Pending":
            for cell in sheet[row]:
                cell.fill = REVIEW_FILL


def _write_latest(sheet, records: dict[str, dict]) -> None:
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for record in sorted(records.values(), key=lambda item: (item["model"], item["retailer"])):
        sheet.append([_iso_cell(record.get(LATEST_FIELD_MAP[header])) for header in LATEST_HEADERS])
        if record.get("review_status") == "Pending":
            for cell in sheet[sheet.max_row]:
                cell.fill = REVIEW_FILL
    sheet.auto_filter.ref = f"A1:P{max(sheet.max_row, 1)}"
    sheet.freeze_panes = "D2"


def _append_changes(sheet, changes: list[dict]) -> None:
    for change in changes:
        sheet.append([_iso_cell(change.get(CHANGE_FIELD_MAP[header])) for header in CHANGE_HEADERS])
        if change.get("review_status") == "Pending":
            for cell in sheet[sheet.max_row]:
                cell.fill = REVIEW_FILL
    sheet.auto_filter.ref = f"A1:O{max(sheet.max_row, 1)}"
    sheet.freeze_panes = "C2"


def _resize_tables(sheet) -> None:
    for table in sheet.tables.values():
        table.ref = f"A1:{sheet.cell(1, sheet.max_column).column_letter}{max(sheet.max_row, 2)}"


def _validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=False)
    expected = {"Latest": LATEST_HEADERS, "Changes": CHANGE_HEADERS, "Run Log": RUN_LOG_HEADERS}
    for sheet_name, headers in expected.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet: {sheet_name}")
        actual = [workbook[sheet_name].cell(1, col).value for col in range(1, len(headers) + 1)]
        if actual != headers:
            raise ValueError(f"Unexpected headers in {sheet_name}")


def apply_run_to_workbook(
    workbook_path: Path,
    results: list[price_check.Result],
    metadata: RunMetadata,
    backup_dir: Path,
) -> WorkbookSummary:
    """Update the Daily workbook through a validated temporary file."""
    workbook_path = Path(workbook_path)
    backup_dir = Path(backup_dir)
    workbook = load_workbook(workbook_path)
    run_sheet = workbook["Run Log"]
    latest_sheet = workbook["Latest"]
    changes_sheet = workbook["Changes"]
    for sheet in (latest_sheet, changes_sheet, run_sheet):
        _compact_data_rows(sheet)
    if metadata.run_id in _existing_values(run_sheet, "Run ID"):
        return WorkbookSummary(metadata.run_id, len(results), 0, 0, 0, 0, 0, "", True)

    previous = _read_latest(latest_sheet)
    event_ids = _existing_values(changes_sheet, "Event ID")
    reconciliation = reconcile_results(previous, results, metadata.run_id, metadata.completed_at, event_ids)
    _write_latest(latest_sheet, reconciliation.latest)
    _append_changes(changes_sheet, reconciliation.changes)
    _paint_pending_rows(changes_sheet)

    ok_count = sum(result.status == "OK" for result in results)
    oos_count = sum(result.status == "OOS" for result in results)
    error_count = len(results) - ok_count - oos_count
    review_count = sum(price_check.is_review_needed(result) for result in results)
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    backup_path = backup_dir / f"{workbook_path.stem}_backup_{stamp}.xlsx"
    shutil.copy2(workbook_path, backup_path)

    run_sheet.append([
        metadata.run_id, metadata.parent_run_id, metadata.run_type, _iso_cell(metadata.started_at),
        _iso_cell(metadata.completed_at), len(results), ok_count, oos_count, error_count,
        len(reconciliation.changes), review_count, "SUCCESS" if error_count == 0 else "PARTIAL",
        metadata.snapshot_path, metadata.review_path, str(backup_path),
    ])
    run_sheet.auto_filter.ref = f"A1:O{run_sheet.max_row}"
    run_sheet.freeze_panes = "D2"
    for sheet in (latest_sheet, changes_sheet, run_sheet):
        _resize_tables(sheet)

    temp_path = workbook_path.with_name(f".{workbook_path.stem}.{metadata.run_id}.tmp.xlsx")
    try:
        workbook.save(temp_path)
        _validate_workbook(temp_path)
        os.replace(temp_path, workbook_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    return WorkbookSummary(
        metadata.run_id, len(results), ok_count, oos_count, error_count,
        len(reconciliation.changes), review_count, str(backup_path), False,
        tuple(reconciliation.changes),
    )


def retry_targets_from_snapshot(snapshot_path: Path) -> list[price_check.Target]:
    targets = []
    with Path(snapshot_path).open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            if row.get("status") not in {"OK", "OOS"} and row.get("url", "").startswith("http"):
                targets.append(price_check.Target(row["model"], row["retailer"], row["url"]))
    return targets


def make_run_id(run_type: str, now: dt.datetime) -> str:
    date_part = now.strftime("%Y%m%d")
    if run_type in {"primary", "retry"}:
        return f"daily-{date_part}-{run_type}"
    return f"daily-{date_part}-manual-{now.strftime('%H%M%S')}"


def write_daily_snapshot(results: list[price_check.Result], snapshot_dir: Path, run_id: str) -> Path:
    snapshot_dir = Path(snapshot_dir)
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    path = snapshot_dir / f"price_snapshot_{run_id}.csv"
    price_check.write_results_csv(results, path)
    return path


def daily_summary_message(
    summary: WorkbookSummary,
    changes: list[dict] | tuple[dict, ...],
    retry_at: str = "11:30",
) -> str:
    lines = [
        f"Price Check Daily {summary.run_id}: {summary.change_count} biến động, "
        f"{summary.ok_count} có giá, {summary.oos_count} OOS, {summary.error_count} lỗi, "
        f"{summary.review_count} cần kiểm tra thủ công."
    ]
    for change in list(changes)[:20]:
        old_value = price_check.format_price(change.get("old_value"))
        new_value = price_check.format_price(change.get("new_value"))
        pending = " [CHECK THỦ CÔNG]" if change.get("review_status") == "Pending" else ""
        lines.append(
            f"- {change['retailer']} · {change['model']}: {old_value} → {new_value} "
            f"({change['change_type']}){pending}"
        )
    if len(changes) > 20:
        lines.append(f"- Còn {len(changes) - 20} biến động trong workbook.")
    if summary.error_count:
        lines.append(f"Có {summary.error_count} lỗi; hệ thống sẽ retry lúc {retry_at}.")
    lines.append(f"Workbook backup: {summary.backup_path or 'không tạo (dry-run/no-op)'}")
    return "\n".join(lines)


def _now() -> dt.datetime:
    return dt.datetime.now(LOCAL_TIMEZONE)


def _fetch_targets(targets: list[price_check.Target], delay: float, overrides: dict) -> list[price_check.Result]:
    results = []
    for index, target in enumerate(targets, start=1):
        print(f"[{index}/{len(targets)}] {target.retailer}: {target.model}")
        results.append(price_check.fetch_target(target, delay, overrides))
    return results


def _primary_snapshot_path(snapshot_dir: Path, now: dt.datetime) -> Path:
    return Path(snapshot_dir) / f"price_snapshot_daily-{now.strftime('%Y%m%d')}-primary.csv"


def _risk_and_backcheck(
    results: list[price_check.Result],
    weekly_workbook: Path,
    render_dir: Path,
    no_back_check: bool,
) -> None:
    source_sheet = price_check.latest_week_sheet(price_check.workbook_sheets(weekly_workbook))
    price_check.add_initial_risk_flags(results, weekly_workbook, source_sheet)
    price_check.run_render_backcheck(results, render_dir, enabled=not no_back_check)


def _dry_run_summary(run_id: str, results: list[price_check.Result]) -> WorkbookSummary:
    ok_count = sum(result.status == "OK" for result in results)
    oos_count = sum(result.status == "OOS" for result in results)
    return WorkbookSummary(
        run_id=run_id,
        total_targets=len(results),
        ok_count=ok_count,
        oos_count=oos_count,
        error_count=len(results) - ok_count - oos_count,
        change_count=0,
        review_count=sum(price_check.is_review_needed(result) for result in results),
        backup_path="",
    )


def run(args: argparse.Namespace) -> int:
    now = _now()
    run_id = make_run_id(args.run_type, now)
    daily_dir = Path(args.daily_dir)
    snapshot_dir = daily_dir / "snapshots"
    review_dir = daily_dir / "reviews"
    render_dir = daily_dir / "render"
    backup_dir = daily_dir / "backups"
    workbook_path = Path(args.workbook)
    link_workbook = Path(args.link_workbook)
    weekly_workbook = Path(args.weekly_workbook)
    overrides = price_check.read_price_overrides(Path(args.override_file))

    parent_run_id = ""
    if args.run_type == "retry":
        primary_snapshot = _primary_snapshot_path(snapshot_dir, now)
        if not primary_snapshot.exists():
            print(f"NO_RETRY: primary snapshot not found: {primary_snapshot}")
            return 2
        targets = retry_targets_from_snapshot(primary_snapshot)
        parent_run_id = make_run_id("primary", now)
        if not targets:
            print(json.dumps({"status": "NO_RETRY", "reason": "No failed targets", "parent_run_id": parent_run_id}))
            return 0
    else:
        targets = price_check.read_link_targets(link_workbook)

    if args.limit:
        targets = targets[: args.limit]
    if not targets:
        print("No targets to process.")
        return 0

    started_at = _now()
    results = _fetch_targets(targets, args.delay, overrides)
    if args.run_type != "retry" and not args.limit:
        source_sheet = price_check.latest_week_sheet(price_check.workbook_sheets(weekly_workbook))
        all_targets = price_check.read_link_targets(link_workbook)
        results.extend(price_check.missing_link_results(weekly_workbook, source_sheet, link_workbook, all_targets))
    _risk_and_backcheck(results, weekly_workbook, render_dir, args.no_back_check)
    snapshot_path = write_daily_snapshot(results, snapshot_dir, run_id)
    review_path = price_check.write_review_report(results, review_dir, run_id)
    completed_at = _now()

    if args.dry_run:
        summary = _dry_run_summary(run_id, results)
    else:
        if not workbook_path.exists():
            raise FileNotFoundError(f"Daily workbook does not exist: {workbook_path}")
        metadata = RunMetadata(
            run_id=run_id,
            parent_run_id=parent_run_id,
            run_type=args.run_type,
            started_at=started_at.isoformat(timespec="seconds"),
            completed_at=completed_at.isoformat(timespec="seconds"),
            snapshot_path=str(snapshot_path),
            review_path=str(review_path or ""),
        )
        summary = apply_run_to_workbook(workbook_path, results, metadata, backup_dir)

    payload = {
        "run_id": summary.run_id,
        "run_type": args.run_type,
        "targets": summary.total_targets,
        "ok": summary.ok_count,
        "oos": summary.oos_count,
        "errors": summary.error_count,
        "changes": summary.change_count,
        "review_needed": summary.review_count,
        "already_recorded": summary.already_recorded,
        "snapshot_path": str(snapshot_path),
        "review_path": str(review_path or ""),
        "workbook_path": str(workbook_path),
        "backup_path": summary.backup_path,
    }
    print("DAILY_REPORT_JSON=" + json.dumps(payload, ensure_ascii=False))
    print(daily_summary_message(summary, summary.changes))
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Run Price Check Daily without modifying Price Check.xlsx.")
    parser.add_argument("--run-type", choices=("primary", "retry", "manual"), default="manual")
    parser.add_argument("--workbook", default=str(DEFAULT_DAILY_WORKBOOK))
    parser.add_argument("--weekly-workbook", default=str(DEFAULT_WEEKLY_WORKBOOK))
    parser.add_argument("--link-workbook", default=str(DEFAULT_LINK_WORKBOOK))
    parser.add_argument("--override-file", default=str(DEFAULT_OVERRIDE_FILE))
    parser.add_argument("--daily-dir", default=str(DEFAULT_DAILY_DIR))
    parser.add_argument("--delay", type=float, default=0.5)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-back-check", action="store_true")
    args = parser.parse_args()
    raise SystemExit(run(args))


if __name__ == "__main__":
    main()
