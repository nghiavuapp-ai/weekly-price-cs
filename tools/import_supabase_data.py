"""Build and upload deterministic Supabase historical-import payloads.

The payload builders are pure: callers provide row dictionaries and receive
JSON-compatible table records. File-system access and network writes live at
the command-line boundary so payload generation can be tested independently.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import ssl
import urllib.parse
import urllib.request
import uuid
from pathlib import Path
from typing import Callable, Iterable, Mapping, Sequence
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
import certifi

from clean_price_data import build_clean_data, canonicalize_model, normalize_price


ROOT = Path(__file__).resolve().parents[1]
LOCAL_TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")
APPLE_WEEK_ANCHOR_DATE = dt.date(2026, 9, 6)
APPLE_WEEK_ANCHOR_ID = "W11Q4FY26"

TABLES = (
    "products",
    "retailers",
    "product_links",
    "price_overrides",
    "crawl_runs",
    "price_observations",
    "price_corrections",
)
UPLOAD_ORDER = (
    "retailers",
    "products",
    "product_links",
    "price_overrides",
    "crawl_runs",
    "price_observations",
    "price_corrections",
)
APPEND_ONLY_TABLES = {"price_observations", "price_corrections"}
REVIEW_STATUS_FLAGS = {
    "WARN_HTML_RENDER_MISMATCH",
    "WARN_RENDER_UNAVAILABLE",
    "WARN_STOCK_ACTION_UNVERIFIED",
    "WARN_OUTLIER",
    "WARN_SUSPICIOUS_LOW",
    "WARN_FETCH_STATUS",
}
RETAILER_ALIASES = {
    "CELL PHONE S": ("CPS", "CPS"),
    "CELLPHONES": ("CPS", "CPS"),
    "CPS": ("CPS", "CPS"),
    "FPT": ("FPT", "FPT"),
    "FPTSHOP": ("FPT", "FPT"),
    "MW": ("MW", "MW"),
    "SD": ("SHOPDUNK", "Shopdunk"),
    "SHOPDUNK": ("SHOPDUNK", "Shopdunk"),
    "TGDD": ("MW", "MW"),
    "THEGIOIDIDONG": ("MW", "MW"),
    "VIETTEL": ("VIETTEL", "Viettel"),
    "VT": ("VIETTEL", "Viettel"),
    "HH": ("HOANGHA", "Hoang Ha"),
    "HOANG HA": ("HOANGHA", "Hoang Ha"),
    "HOÀNG HÀ": ("HOANGHA", "Hoang Ha"),
    "HOANGHA": ("HOANGHA", "Hoang Ha"),
}
TABLE_SORT_KEYS = {
    "products": lambda row: (row["name"], row["id"]),
    "retailers": lambda row: (row["code"], row["id"]),
    "product_links": lambda row: (row["product_id"], row["url"], row["id"]),
    "price_overrides": lambda row: (row["product_link_id"], row["id"]),
    "crawl_runs": lambda row: (row["period_start"], row["run_key"], row["id"]),
    "price_observations": lambda row: (
        row["period_start"], row["period_key"], row["product_id"], row["id"]
    ),
    "price_corrections": lambda row: (row.get("created_at", ""), row["id"]),
}


def _empty_payload() -> dict[str, list[dict]]:
    return {table: [] for table in TABLES}


def _stable_id(kind: str, *parts: object) -> str:
    value = "/".join(("price-check", kind, *(str(part) for part in parts)))
    return str(uuid.uuid5(uuid.NAMESPACE_URL, value))


def _clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ")).strip()


def _retailer_record(value: object) -> dict:
    key = _clean_text(value).upper()
    if key not in RETAILER_ALIASES:
        raise ValueError(f"Unsupported retailer: {_clean_text(value) or '<blank>'}")
    code, name = RETAILER_ALIASES[key]
    return {"id": _stable_id("retailer", code), "code": code, "name": name, "active": True}


def _product_record(
    value: object,
    *,
    category: object | None = None,
    capacity: object | None = None,
) -> dict:
    identity = canonicalize_model(_clean_text(value))
    name = identity["canonical_model"]
    if not name:
        raise ValueError("Product model cannot be blank")
    return {
        "id": _stable_id("product", name),
        "name": name,
        "category": _clean_text(category) or identity["category"],
        "capacity": _clean_text(capacity) or identity["capacity"],
        "active": True,
    }


def _link_record(product: Mapping, retailer: Mapping, url: object) -> dict:
    clean_url = _clean_text(url)
    if not clean_url.startswith(("https://", "http://")):
        raise ValueError(f"Invalid product URL for {product['name']} / {retailer['name']}")
    return {
        "id": _stable_id("product-link", product["id"], retailer["id"], clean_url),
        "product_id": product["id"],
        "retailer_id": retailer["id"],
        "url": clean_url,
        "active": True,
    }


def _sorted_payload(payload: Mapping[str, Sequence[dict]]) -> dict[str, list[dict]]:
    return {
        table: sorted((dict(row) for row in payload.get(table, [])), key=TABLE_SORT_KEYS[table])
        for table in TABLES
    }


def build_catalog_payload(
    link_rows: Iterable[Mapping[str, object]],
    override_rows: Iterable[Mapping[str, object]] = (),
) -> dict[str, list[dict]]:
    """Build deterministic configuration rows from link and override records."""
    payload = _empty_payload()
    products: dict[str, dict] = {}
    retailers: dict[str, dict] = {}
    links: dict[str, dict] = {}
    links_by_url: dict[str, dict] = {}

    for source in link_rows:
        model_key = next((key for key in source if _clean_text(key).casefold() == "model"), None)
        if model_key is None or not _clean_text(source.get(model_key)):
            continue
        product = _product_record(source[model_key])
        products[product["id"]] = product
        for column, value in source.items():
            if column == model_key or not _clean_text(value):
                continue
            retailer = _retailer_record(column)
            link = _link_record(product, retailer, value)
            retailers[retailer["id"]] = retailer
            links[link["id"]] = link
            links_by_url[link["url"]] = link

    overrides: dict[str, dict] = {}
    for source in override_rows:
        url = _clean_text(source.get("url"))
        if not url:
            continue
        link = links_by_url.get(url)
        if link is None:
            raise ValueError(f"Override URL is not present in the product-link catalog: {url}")
        price = normalize_price(source.get("value"))
        if price is None:
            raise ValueError(f"Override price must be a positive integer for: {url}")
        record = {
            "id": _stable_id("price-override", link["id"]),
            "product_link_id": link["id"],
            "override_price": price,
            "note": _clean_text(source.get("note")),
            "active": True,
        }
        overrides[record["id"]] = record

    payload["products"] = list(products.values())
    payload["retailers"] = list(retailers.values())
    payload["product_links"] = list(links.values())
    payload["price_overrides"] = list(overrides.values())
    return _sorted_payload(payload)


def _week_serial(week_id: str) -> int:
    match = re.fullmatch(r"W(\d{1,2})Q([1-4])FY(\d+)", _clean_text(week_id), re.IGNORECASE)
    if not match:
        raise ValueError(f"Invalid Apple fiscal week: {week_id}")
    week, quarter, fiscal_year = (int(value) for value in match.groups())
    if week < 1 or week > 13:
        raise ValueError(f"Invalid Apple fiscal week: {week_id}")
    return fiscal_year * 52 + (quarter - 1) * 13 + (week - 1)


def _week_start(week_id: str) -> dt.date:
    delta = _week_serial(week_id) - _week_serial(APPLE_WEEK_ANCHOR_ID)
    return APPLE_WEEK_ANCHOR_DATE + dt.timedelta(weeks=delta)


def _local_iso(value: object) -> str:
    if isinstance(value, dt.datetime):
        parsed = value
    elif isinstance(value, dt.date):
        parsed = dt.datetime.combine(value, dt.time())
    else:
        text = _clean_text(value)
        if text.endswith("Z"):
            text = text[:-1] + "+00:00"
        parsed = dt.datetime.fromisoformat(text)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=LOCAL_TIMEZONE)
    return parsed.isoformat(timespec="seconds")


def _risk_flags(value: object) -> list[str]:
    if isinstance(value, (list, tuple, set)):
        candidates = value
    else:
        candidates = str(value or "").split("|")
    return sorted({_clean_text(flag) for flag in candidates if _clean_text(flag)})


def _observation_id(run_id: str, product_id: str, retailer_id: str) -> str:
    return _stable_id("observation", run_id, product_id, retailer_id)


def build_weekly_payload(rows: Iterable[Mapping[str, object]]) -> dict[str, list[dict]]:
    """Convert normalized Weekly workbook rows into immutable run observations."""
    grouped: dict[str, list[Mapping[str, object]]] = {}
    for row in rows:
        week_id = _clean_text(row.get("week_id"))
        if week_id:
            grouped.setdefault(week_id, []).append(row)

    payload = _empty_payload()
    products: dict[str, dict] = {}
    retailers: dict[str, dict] = {}
    observations: dict[str, dict] = {}

    for week_id, week_rows in grouped.items():
        period_start = _week_start(week_id)
        observed_at = dt.datetime.combine(period_start, dt.time(12), LOCAL_TIMEZONE).isoformat(
            timespec="seconds"
        )
        run_key = f"historical-weekly:{week_id}"
        run_id = _stable_id("crawl-run", run_key)
        week_observations: dict[str, dict] = {}

        for source in week_rows:
            product = _product_record(
                source.get("canonical_model") or source.get("raw_model"),
                category=source.get("category"),
                capacity=source.get("capacity"),
            )
            retailer = _retailer_record(source.get("partner"))
            products[product["id"]] = product
            retailers[retailer["id"]] = retailer
            price = normalize_price(source.get("price_vnd"))
            stock = _clean_text(source.get("stock_status")).casefold()
            in_stock = True if price is not None else False if stock == "oos" else None
            fetch_ok = stock != "missing"
            risks = []
            if not fetch_ok:
                risks.append("HISTORICAL_MISSING")
            if _clean_text(source.get("mapping_status")).casefold() == "review":
                risks.append("HISTORICAL_MAPPING_REVIEW")
            review_status = "pending" if risks else "confirmed"
            record_id = _observation_id(run_id, product["id"], retailer["id"])
            week_observations[record_id] = {
                "id": record_id,
                "run_id": run_id,
                "product_id": product["id"],
                "retailer_id": retailer["id"],
                "product_link_id": None,
                "period_type": "weekly",
                "period_key": week_id,
                "period_start": period_start.isoformat(),
                "observed_at": observed_at,
                "price_vnd": price,
                "in_stock": in_stock,
                "fetch_ok": fetch_ok,
                "review_status": review_status,
                "confidence": "historical",
                "risk_flags": risks,
                "evidence": {
                    "mapping_status": source.get("mapping_status"),
                    "price_change_note": source.get("price_change_note"),
                    "raw_model": source.get("raw_model"),
                    "source_row": source.get("source_row"),
                    "source_sheet": source.get("source_sheet"),
                },
                "source_method": "historical_workbook",
                "error_message": None,
            }

        week_values = list(week_observations.values())
        payload["crawl_runs"].append({
            "id": run_id,
            "run_key": run_key,
            "parent_run_id": None,
            "run_type": "historical_weekly",
            "period_type": "weekly",
            "period_key": week_id,
            "period_start": period_start.isoformat(),
            "scheduled_for": observed_at,
            "started_at": observed_at,
            "completed_at": observed_at,
            "status": "success",
            "total_targets": len(week_values),
            "ok_count": sum(row["in_stock"] is True for row in week_values),
            "oos_count": sum(row["in_stock"] is False for row in week_values),
            "error_count": sum(not row["fetch_ok"] for row in week_values),
            "change_count": 0,
            "review_count": sum(row["review_status"] == "pending" for row in week_values),
            "source": "historical_import",
        })
        observations.update(week_observations)

    payload["products"] = list(products.values())
    payload["retailers"] = list(retailers.values())
    payload["price_observations"] = list(observations.values())
    return _sorted_payload(payload)


def _integer(value: object, default: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _daily_evidence(source: Mapping[str, object]) -> dict:
    fields = (
        "value",
        "html_price",
        "rendered_price",
        "previous_week_price",
        "base_value",
        "discount",
        "promo_note",
        "availability",
        "availability_method",
        "availability_text",
        "purchase_action",
        "purchase_action_method",
        "purchase_action_text",
        "decision_note",
    )
    return {field: source[field] for field in fields if source.get(field) not in (None, "")}


def build_daily_snapshot_payload(
    run: Mapping[str, object],
    rows: Iterable[Mapping[str, object]],
) -> dict[str, list[dict]]:
    """Convert one logged Daily snapshot into a run and raw observations."""
    run_key = _clean_text(run.get("run_id"))
    if not run_key:
        raise ValueError("Daily run_id cannot be blank")
    run_id = _stable_id("crawl-run", run_key)
    source_rows = list(rows)
    fallback_time = run.get("completed_at") or run.get("started_at")
    if not fallback_time:
        fallback_time = next((row.get("fetched_at") for row in source_rows if row.get("fetched_at")), None)
    if not fallback_time:
        raise ValueError(f"Daily run has no timestamp: {run_key}")
    completed_at = _local_iso(run.get("completed_at") or fallback_time)
    started_at = _local_iso(run.get("started_at") or fallback_time)
    period_start = dt.datetime.fromisoformat(completed_at).date().isoformat()

    payload = _empty_payload()
    products: dict[str, dict] = {}
    retailers: dict[str, dict] = {}
    links: dict[str, dict] = {}
    observations: dict[str, dict] = {}

    for source in source_rows:
        if not _clean_text(source.get("model")) or not _clean_text(source.get("retailer")):
            continue
        product = _product_record(source.get("model"))
        retailer = _retailer_record(source.get("retailer"))
        products[product["id"]] = product
        retailers[retailer["id"]] = retailer
        link = None
        if _clean_text(source.get("url")):
            link = _link_record(product, retailer, source.get("url"))
            links[link["id"]] = link

        status = _clean_text(source.get("status")).upper()
        price = normalize_price(source.get("final_price")) or normalize_price(source.get("value"))
        fetch_ok = status in {"OK", "OOS"}
        in_stock = True if status == "OK" and price is not None else False if status == "OOS" else None
        risks = _risk_flags(source.get("risk_flags"))
        confidence = _clean_text(source.get("confidence")).upper()
        verified_oos = status == "OOS" and confidence in {"VERIFIED_OOS", "OOS_CONFIRMED"}
        requires_review = not verified_oos and (
            not fetch_ok
            or confidence == "REVIEW"
            or bool(set(risks) & REVIEW_STATUS_FLAGS)
        )
        record_id = _observation_id(run_id, product["id"], retailer["id"])
        observations[record_id] = {
            "id": record_id,
            "run_id": run_id,
            "product_id": product["id"],
            "retailer_id": retailer["id"],
            "product_link_id": link["id"] if link else None,
            "period_type": "daily",
            "period_key": period_start,
            "period_start": period_start,
            "observed_at": _local_iso(source.get("fetched_at") or fallback_time),
            "price_vnd": price,
            "in_stock": in_stock,
            "fetch_ok": fetch_ok,
            "review_status": "pending" if requires_review else "confirmed",
            "confidence": _clean_text(source.get("confidence")) or None,
            "risk_flags": risks,
            "evidence": _daily_evidence(source),
            "source_method": _clean_text(source.get("source_method")) or None,
            "error_message": _clean_text(source.get("error")) or None,
        }

    values = list(observations.values())
    raw_status = _clean_text(run.get("status")).casefold().replace(" ", "_")
    status = {
        "success": "success",
        "partial": "partial",
        "failed": "failed",
        "running": "running",
        "no_retry": "no_retry",
    }.get(raw_status, "partial" if any(not row["fetch_ok"] for row in values) else "success")
    parent_run_key = _clean_text(run.get("parent_run_id"))
    payload["crawl_runs"] = [{
        "id": run_id,
        "run_key": run_key,
        "parent_run_id": _stable_id("crawl-run", parent_run_key) if parent_run_key else None,
        "run_type": _clean_text(run.get("run_type")).casefold() or "manual",
        "period_type": "daily",
        "period_key": period_start,
        "period_start": period_start,
        "scheduled_for": started_at,
        "started_at": started_at,
        "completed_at": completed_at,
        "status": status,
        "total_targets": _integer(run.get("total_targets"), len(values)),
        "ok_count": _integer(run.get("ok_count"), sum(row["in_stock"] is True for row in values)),
        "oos_count": _integer(run.get("oos_count"), sum(row["in_stock"] is False for row in values)),
        "error_count": _integer(run.get("error_count"), sum(not row["fetch_ok"] for row in values)),
        "change_count": _integer(run.get("change_count"), 0),
        "review_count": _integer(
            run.get("review_count"), sum(row["review_status"] == "pending" for row in values)
        ),
        "source": "daily_import",
    }]
    payload["products"] = list(products.values())
    payload["retailers"] = list(retailers.values())
    payload["product_links"] = list(links.values())
    payload["price_observations"] = values
    return _sorted_payload(payload)


def merge_payloads(
    parts: Iterable[Mapping[str, Sequence[Mapping]]],
    generated_at: str,
) -> dict[str, object]:
    """Merge table payloads by deterministic ID and reject conflicting duplicates."""
    merged: dict[str, dict[str, dict]] = {table: {} for table in TABLES}
    for part in parts:
        for table in TABLES:
            for source in part.get(table, []):
                row = dict(source)
                row_id = _clean_text(row.get("id"))
                if not row_id:
                    raise ValueError(f"{table} row is missing an id")
                existing = merged[table].get(row_id)
                if existing is not None and existing != row:
                    raise ValueError(f"Conflicting deterministic {table} row: {row_id}")
                merged[table][row_id] = row
    result: dict[str, object] = {"generated_at": generated_at}
    result.update(_sorted_payload({table: list(rows.values()) for table, rows in merged.items()}))
    return result


def _workbook_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [_clean_text(value) for value in next(values, ())]
        return [dict(zip(headers, row)) for row in values if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def _csv_rows(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _daily_batches(workbook_path: Path) -> list[tuple[dict, list[dict]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Run Log" not in workbook.sheetnames:
            return []
        sheet = workbook["Run Log"]
        headers = [_clean_text(cell.value) for cell in sheet[1]]
        field_map = {
            "Run ID": "run_id",
            "Parent Run ID": "parent_run_id",
            "Run Type": "run_type",
            "Started At": "started_at",
            "Completed At": "completed_at",
            "Total Targets": "total_targets",
            "OK": "ok_count",
            "OOS": "oos_count",
            "Errors": "error_count",
            "Changes": "change_count",
            "Review Needed": "review_count",
            "Status": "status",
            "Snapshot Path": "snapshot_path",
        }
        batches = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            source = dict(zip(headers, values))
            if not source.get("Run ID") or not source.get("Snapshot Path"):
                continue
            run = {target: source.get(header) for header, target in field_map.items()}
            snapshot_path = Path(str(run.pop("snapshot_path")))
            if not snapshot_path.is_absolute():
                snapshot_path = workbook_path.parent / snapshot_path
            if not snapshot_path.is_file():
                raise FileNotFoundError(f"Logged Daily snapshot does not exist: {snapshot_path}")
            batches.append((run, _csv_rows(snapshot_path)))
        return batches
    finally:
        workbook.close()


def build_project_payload(root: Path, generated_at: str) -> dict[str, object]:
    """Read the project sources and assemble one deterministic import document."""
    root = Path(root)
    catalog = build_catalog_payload(
        _workbook_rows(root / "Link Product.xlsx"),
        _csv_rows(root / "Price Overrides.csv") if (root / "Price Overrides.csv").is_file() else (),
    )
    weekly_rows, _, _ = build_clean_data(root / "Price Check.xlsx")
    parts: list[Mapping[str, Sequence[Mapping]]] = [catalog, build_weekly_payload(weekly_rows)]
    for run, rows in _daily_batches(root / "Price Check Daily.xlsx"):
        parts.append(build_daily_snapshot_payload(run, rows))
    return merge_payloads(parts, generated_at=generated_at)


def _default_sender(request: urllib.request.Request, timeout: int) -> None:
    context = ssl.create_default_context(cafile=certifi.where())
    with urllib.request.urlopen(request, timeout=timeout, context=context):  # noqa: S310 - explicit CLI destination.
        return None


def upsert_payload(
    payload: Mapping[str, object],
    supabase_url: str,
    service_key: str,
    batch_size: int = 500,
    sender: Callable[[urllib.request.Request, int], object] | None = None,
) -> dict[str, int]:
    """Upload batches through authenticated PostgREST upserts without logging credentials."""
    if not _clean_text(supabase_url) or not service_key:
        raise ValueError("Supabase URL and service key are required")
    if batch_size < 1:
        raise ValueError("batch_size must be at least 1")
    send = sender or _default_sender
    base_url = supabase_url.rstrip("/")
    counts: dict[str, int] = {}
    for table in UPLOAD_ORDER:
        rows = list(payload.get(table, []))
        counts[table] = len(rows)
        resolution = "ignore-duplicates" if table in APPEND_ONLY_TABLES else "merge-duplicates"
        for start in range(0, len(rows), batch_size):
            batch = rows[start : start + batch_size]
            query = urllib.parse.urlencode({"on_conflict": "id"})
            request = urllib.request.Request(
                f"{base_url}/rest/v1/{table}?{query}",
                data=json.dumps(batch, ensure_ascii=False, separators=(",", ":")).encode("utf-8"),
                headers={
                    "apikey": service_key,
                    "Authorization": f"Bearer {service_key}",
                    "Content-Type": "application/json",
                    "Prefer": f"resolution={resolution},return=minimal",
                },
                method="POST",
            )
            send(request, 60)
    return counts


def _generated_at() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT, help="Price Check project root")
    parser.add_argument("--generated-at", default=None, help="Explicit ISO timestamp for reproducible dry runs")
    parser.add_argument("--batch-size", type=int, default=500)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="Print import JSON without network writes")
    mode.add_argument("--apply", action="store_true", help="Upsert the payload into Supabase")
    args = parser.parse_args(argv)

    payload = build_project_payload(args.root, generated_at=args.generated_at or _generated_at())
    if not args.apply:
        print(json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
        return 0

    supabase_url = os.environ.get("SUPABASE_URL", "")
    service_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
    if not supabase_url or not service_key:
        parser.error("--apply requires SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY")
    counts = upsert_payload(payload, supabase_url, service_key, batch_size=args.batch_size)
    print(json.dumps({"status": "uploaded", "counts": counts}, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
