"""Normalize Price Check.xlsx into browser-ready data for the weekly retail-price dashboard."""
import json
import os
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Price Check.xlsx"
OUTPUT = ROOT / "dashboard-data.js"
PARTNER_ALIASES = {
    "MW": "MW", "FPT": "FPT", "Viettel": "Viettel", "CPS": "CPS", "Shopdunk": "Shopdunk",
    "SD": "Shopdunk", "TZ": "TZ", "Hoang Ha": "Hoang Ha", "HH": "Hoang Ha", "VT": "Viettel",
}

def canonical_week(value):
    match = re.search(r"W(\d+)Q(\d)FY(\d+)", str(value).replace(" ", ""), re.I)
    if not match:
        return None
    week, quarter, year = map(int, match.groups())
    return {"id": f"W{week}Q{quarter}FY{year}", "week": week, "quarter": quarter, "year": year}

def numeric_price(value):
    if isinstance(value, (int, float)) and value > 0:
        return int(value)
    if isinstance(value, str):
        compact = value.replace(".", "").replace(",", "").strip()
        if compact.isdigit() and int(compact) > 0:
            return int(compact)
    return None

book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
points = []
for sheet in book.worksheets:
    meta = canonical_week(sheet.title)
    if not meta:
        continue
    headers = {}
    for col in range(2, min(sheet.max_column, 12) + 1):
        partner = PARTNER_ALIASES.get(str(sheet.cell(1, col).value).strip()) if sheet.cell(1, col).value else None
        if partner and partner not in headers.values():
            headers[col] = partner
    for row in range(2, sheet.max_row + 1):
        model = sheet.cell(row, 1).value
        if not isinstance(model, str) or not model.strip().lower().startswith("iphone"):
            continue
        model = re.sub(r"\s+", " ", model).strip()
        for col, partner in headers.items():
            raw = sheet.cell(row, col).value
            price = numeric_price(raw)
            if price or str(raw).strip().upper() == "OOS":
                points.append({**meta, "model": model, "partner": partner, "price": price, "oos": price is None})

points.sort(key=lambda p: (p["year"], p["quarter"], p["week"], p["model"], p["partner"]))
payload = {"source": SOURCE.name, "generatedAt": __import__("datetime").datetime.now().strftime("%Y-%m-%d %H:%M"), "points": points}
OUTPUT.write_text("window.PRICE_DATA = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(f"Wrote {len(points)} price records to {OUTPUT}")
