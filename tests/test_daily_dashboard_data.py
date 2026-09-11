import csv
import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

from daily_dashboard_data import apple_week_id_for_date, load_daily_payload


RUN_LOG_HEADERS = [
    "Run ID", "Parent Run ID", "Run Type", "Started At", "Completed At",
    "Total Targets", "OK", "OOS", "Errors", "Changes", "Review Needed",
    "Status", "Snapshot Path", "Review Path", "Backup Path",
]
SNAPSHOT_HEADERS = [
    "model", "retailer", "value", "final_price", "html_price", "rendered_price",
    "previous_week_price", "base_value", "discount", "promo_note", "source_method",
    "availability", "availability_method", "availability_text", "risk_flags",
    "confidence", "decision_note", "status", "error", "fetched_at", "url",
]


class AppleWeekTests(unittest.TestCase):
    def test_maps_sunday_through_saturday_to_same_apple_week(self):
        self.assertEqual(apple_week_id_for_date(dt.date(2026, 9, 6)), "W11Q4FY26")
        self.assertEqual(apple_week_id_for_date(dt.date(2026, 9, 12)), "W11Q4FY26")

    def test_rolls_week_thirteen_to_next_fiscal_year(self):
        self.assertEqual(apple_week_id_for_date(dt.date(2026, 9, 26)), "W13Q4FY26")
        self.assertEqual(apple_week_id_for_date(dt.date(2026, 9, 27)), "W1Q1FY27")


class DailyPayloadTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.workbook_path = self.root / "Price Check Daily.xlsx"
        workbook = Workbook()
        run_log = workbook.active
        run_log.title = "Run Log"
        run_log.append(RUN_LOG_HEADERS)
        workbook.save(self.workbook_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def write_snapshot(self, name, rows):
        path = self.root / name
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=SNAPSHOT_HEADERS)
            writer.writeheader()
            for row in rows:
                writer.writerow({field: row.get(field, "") for field in SNAPSHOT_HEADERS})
        return path

    def log_run(self, run_id, completed_at, snapshot_path, run_type="primary"):
        from openpyxl import load_workbook

        workbook = load_workbook(self.workbook_path)
        workbook["Run Log"].append([
            run_id, "", run_type, completed_at, completed_at, 1, 1, 0, 0, 0, 0,
            "SUCCESS", str(snapshot_path), "", "",
        ])
        workbook.save(self.workbook_path)

    def test_uses_only_logged_snapshots_and_normalizes_model(self):
        logged = self.write_snapshot("logged.csv", [{
            "model": "AW S11 (42mm CEL)", "retailer": "FPT", "value": "14090000",
            "final_price": "14090000", "status": "OK", "confidence": "HTML_OK",
            "fetched_at": "2026-09-10T11:00:00",
        }])
        self.write_snapshot("unlogged.csv", [{
            "model": "iPhone 17e 256GB", "retailer": "MW", "value": "1",
            "final_price": "1", "status": "OK", "fetched_at": "2026-09-10T12:00:00",
        }])
        self.log_run("daily-20260910-primary", dt.datetime(2026, 9, 10, 11), logged)

        payload = load_daily_payload(self.workbook_path)

        self.assertEqual(payload["dailyDates"], ["2026-09-10"])
        self.assertEqual(payload["dailyWeeks"], ["W11Q4FY26"])
        self.assertEqual(len(payload["dailyRows"]), 1)
        self.assertEqual(payload["dailyRows"][0]["canonical_model"], "Apple Watch S11 (42mm CEL)")
        self.assertNotIn("url", payload["dailyRows"][0])

    def test_retry_wins_and_failed_checks_carry_forward_as_stale(self):
        day_one = self.write_snapshot("day-one.csv", [
            {"model": "iPhone 15 128GB", "retailer": "FPT", "final_price": "18000000", "status": "OK", "confidence": "HTML_OK", "fetched_at": "2026-09-10T11:00:00"},
            {"model": "iPhone 15 128GB", "retailer": "CPS", "value": "OOS", "status": "OOS", "confidence": "OOS_CONFIRMED", "fetched_at": "2026-09-10T11:01:00"},
        ])
        day_two_primary = self.write_snapshot("day-two-primary.csv", [
            {"model": "iPhone 15 128GB", "retailer": "FPT", "status": "URL_ERROR", "fetched_at": "2026-09-11T11:00:00"},
            {"model": "iPhone 15 128GB", "retailer": "CPS", "status": "URL_ERROR", "fetched_at": "2026-09-11T11:01:00"},
        ])
        day_two_retry = self.write_snapshot("day-two-retry.csv", [
            {"model": "iPhone 15 128GB", "retailer": "FPT", "final_price": "17900000", "status": "OK", "confidence": "VERIFIED_RENDER", "risk_flags": "WARN_WEEK_CHANGE", "fetched_at": "2026-09-11T11:30:00"},
        ])
        self.log_run("daily-20260910-primary", dt.datetime(2026, 9, 10, 11, 5), day_one)
        self.log_run("daily-20260911-primary", dt.datetime(2026, 9, 11, 11, 5), day_two_primary)
        self.log_run("daily-20260911-retry", dt.datetime(2026, 9, 11, 11, 35), day_two_retry, "retry")

        payload = load_daily_payload(self.workbook_path)
        rows = {(row["date"], row["partner"]): row for row in payload["dailyRows"]}

        self.assertEqual(rows[("2026-09-11", "FPT")]["price_vnd"], 17_900_000)
        self.assertFalse(rows[("2026-09-11", "FPT")]["stale"])
        self.assertEqual(rows[("2026-09-11", "CPS")]["stock_status"], "OOS")
        self.assertTrue(rows[("2026-09-11", "CPS")]["stale"])
        self.assertEqual(rows[("2026-09-11", "CPS")]["observed_at"], "2026-09-10T11:01:00")

    def test_keeps_only_dates_in_latest_thirteen_apple_weeks(self):
        for offset in (0, 7, 84, 91):
            day = dt.date(2026, 6, 14) + dt.timedelta(days=offset)
            snapshot = self.write_snapshot(f"snapshot-{day.isoformat()}.csv", [{
                "model": "iPhone 15 128GB", "retailer": "FPT", "final_price": "18000000",
                "status": "OK", "confidence": "HTML_OK", "fetched_at": f"{day.isoformat()}T11:00:00",
            }])
            self.log_run(f"daily-{day.strftime('%Y%m%d')}-primary", dt.datetime.combine(day, dt.time(11, 5)), snapshot)

        payload = load_daily_payload(self.workbook_path, retention_weeks=13)

        self.assertEqual(payload["dailyDates"], ["2026-06-21", "2026-09-06", "2026-09-13"])
        self.assertNotIn("2026-06-14", payload["dailyDates"])


if __name__ == "__main__":
    unittest.main()
