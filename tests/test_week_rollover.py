import importlib.util
import json
import os
import tempfile
import subprocess
import sys
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook


MODULE_PATH = Path(__file__).parents[1] / "tools" / "price_check_tool.py"
SPEC = importlib.util.spec_from_file_location("price_check_tool", MODULE_PATH)
price_check_tool = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = price_check_tool
SPEC.loader.exec_module(price_check_tool)


class WeekRolloverTests(unittest.TestCase):
    def test_rolls_week_13_into_next_quarter(self):
        self.assertEqual(price_check_tool.next_week_name("W13Q3FY26"), "W1Q4FY26")

    def test_rolls_fiscal_year_after_quarter_4(self):
        self.assertEqual(price_check_tool.next_week_name("W13Q4FY26"), "W1Q1FY27")

    def test_copied_formulas_reference_source_week(self):
        workbook = Workbook()
        old = workbook.active
        old.title = "W12Q3FY26"
        source = workbook.copy_worksheet(old)
        source.title = "W13Q3FY26"
        source["M2"] = '=IFERROR(C2-W12Q3FY26!C2,"OOS")'

        target, _ = price_check_tool.prepare_target_sheet(workbook, "W13Q3FY26", "W1Q4FY26")

        self.assertEqual(target["M2"].value, '=IFERROR(C2-\'W13Q3FY26\'!C2,"OOS")')
        self.assertEqual(price_check_tool.count_formula_sheet_refs(target, "W12Q3FY26"), 0)
        self.assertEqual(price_check_tool.count_formula_sheet_refs(target, "W13Q3FY26"), 1)

    def test_existing_target_repairs_missing_formula(self):
        workbook = Workbook()
        old = workbook.active
        old.title = "W12Q3FY26"
        source = workbook.create_sheet("W13Q3FY26")
        source["M7"] = '=IFERROR(C7-W12Q3FY26!C7,"OOS")'
        source["N7"] = '=IF(M7=0,"","changed")'
        target = workbook.create_sheet("W1Q4FY26")

        repaired, _ = price_check_tool.prepare_target_sheet(workbook, "W13Q3FY26", "W1Q4FY26")

        self.assertEqual(repaired["M7"].value, '=IFERROR(C7-\'W13Q3FY26\'!C7,"OOS")')
        self.assertEqual(repaired["N7"].value, '=IF(M7=0,"","changed")')


class AvailabilityTests(unittest.TestCase):
    def test_fpt_public_api_returns_buyable_default_sku(self):
        target = price_check_tool.Target(
            "iPhone 17e 256GB",
            "FPT",
            "https://fptshop.com.vn/dien-thoai/iphone-17e",
        )
        responses = [
            {
                "status": 200,
                "data": {
                    "skus": [
                        {
                            "code": "00926263",
                            "displayName": "iPhone 17e 256GB",
                            "price": 21_490_000,
                            "inventory": 6,
                            "isDefaultSku": True,
                        }
                    ]
                },
            },
            {"status": 200, "data": {"buttonCode": "ORDER", "statusOnWeb": ""}},
        ]

        with mock.patch.object(price_check_tool, "request_json", side_effect=responses):
            result = price_check_tool.fetch_fpt_api_result(target)

        self.assertEqual(result.status, "OK")
        self.assertEqual(result.value, 21_490_000)
        self.assertEqual(result.source_method, "fpt_public_api")
        self.assertEqual(result.purchase_action, "IN_STOCK")
        self.assertEqual(result.purchase_action_text, "ORDER")
        price_check_tool.add_initial_risk_flags_from_previous(
            [result], {(result.model, "FPT"): 21_490_000}
        )
        self.assertNotIn("WARN_STOCK_ACTION_UNVERIFIED", result.risk_flags)
        self.assertFalse(price_check_tool.is_render_candidate(result))

    def test_fpt_public_api_does_not_keep_price_when_status_is_oos(self):
        target = price_check_tool.Target(
            "iPhone 14 128GB",
            "FPT",
            "https://fptshop.com.vn/dien-thoai/iphone-14",
        )
        responses = [
            {
                "status": 200,
                "data": {
                    "skus": [
                        {
                            "code": "00832860",
                            "displayName": "iPhone 14 128GB",
                            "price": 13_990_000,
                            "inventory": 1,
                            "isDefaultSku": True,
                        }
                    ]
                },
            },
            {
                "status": 200,
                "data": {
                    "buttonCode": "REGISTER_IN_ADVANCE",
                    "statusOnWeb": "tam_het_hang",
                },
            },
        ]

        with mock.patch.object(price_check_tool, "request_json", side_effect=responses):
            result = price_check_tool.fetch_fpt_api_result(target)

        self.assertEqual(result.status, "OOS")
        self.assertEqual(result.value, "OOS")
        self.assertEqual(result.confidence, "VERIFIED_OOS")
        self.assertEqual(result.source_method, "fpt_public_api_status")
        self.assertFalse(price_check_tool.is_render_candidate(result))

    def test_fpt_403_falls_back_to_public_api(self):
        target = price_check_tool.Target(
            "iPhone 17e 256GB",
            "FPT",
            "https://fptshop.com.vn/dien-thoai/iphone-17e",
        )
        fallback = price_check_tool.Result(
            target.model,
            target.retailer,
            target.url,
            21_490_000,
            "OK",
            source_method="fpt_public_api",
        )
        error = price_check_tool.HTTPError(target.url, 403, "Forbidden", {}, None)

        with mock.patch.object(price_check_tool, "request_html", side_effect=error), \
             mock.patch.object(price_check_tool, "fetch_fpt_api_result", return_value=fallback) as api:
            result = price_check_tool.fetch_target(target, 0)

        self.assertIs(result, fallback)
        api.assert_called_once_with(target)

    def test_fpt_request_uses_configured_cloud_proxy(self):
        class Headers:
            @staticmethod
            def get_content_charset():
                return "utf-8"

        class Response:
            headers = Headers()

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            @staticmethod
            def read():
                return b"<html>FPT</html>"

        opener = mock.Mock()
        opener.open.return_value = Response()
        with mock.patch.dict(os.environ, {"FPT_PROXY_URL": "http://proxy.test:8000"}), \
             mock.patch.object(price_check_tool.urllib.request, "build_opener", return_value=opener) as build_opener, \
             mock.patch.object(price_check_tool, "urlopen") as direct_open:
            result = price_check_tool.request_html("https://fptshop.com.vn/dien-thoai/iphone-15")

        self.assertEqual(result, "<html>FPT</html>")
        direct_open.assert_not_called()
        self.assertTrue(any(isinstance(handler, price_check_tool.urllib.request.ProxyHandler)
                            for handler in build_opener.call_args.args))

    def test_cloud_relay_is_preferred_for_blocked_retailer_html(self):
        relay_payload = json.dumps({"body": "<html>MW</html>", "status": 200}).encode("utf-8")

        class Headers:
            @staticmethod
            def get_content_charset():
                return "utf-8"

        class Response:
            headers = Headers()

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            @staticmethod
            def read():
                return relay_payload

        with mock.patch.dict(os.environ, {
            "PRICE_FETCH_RELAY_URL": "https://relay.test/api/retailer-fetch",
            "PRICE_FETCH_RELAY_TOKEN": "secret",
        }), mock.patch.object(price_check_tool, "urlopen", return_value=Response()) as open_mock:
            result = price_check_tool.request_html(
                "https://www.thegioididong.com/dtdd/iphone-17-pro-max"
            )

        self.assertEqual(result, "<html>MW</html>")
        request = open_mock.call_args.args[0]
        self.assertEqual(request.full_url, "https://relay.test/api/retailer-fetch")
        self.assertEqual(request.get_header("Authorization"), "Bearer secret")

    def test_cloud_relay_returns_fpt_json_payload(self):
        relay_payload = json.dumps({
            "body": '{"data":{"skus":[]}}',
            "status": 200,
        }).encode("utf-8")

        class Headers:
            @staticmethod
            def get_content_charset():
                return "utf-8"

        class Response:
            headers = Headers()

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return False

            @staticmethod
            def read():
                return relay_payload

        with mock.patch.dict(os.environ, {
            "PRICE_FETCH_RELAY_URL": "https://relay.test/api/retailer-fetch",
            "PRICE_FETCH_RELAY_TOKEN": "secret",
        }), mock.patch.object(price_check_tool, "urlopen", return_value=Response()):
            result = price_check_tool.request_json(
                "https://papi.fptshop.com.vn/gw/v1/public/bff-before-order/product/variant?slug=iphone"
            )

        self.assertEqual(result, {"data": {"skus": []}})

    def test_price_parser_handles_display_and_decimal_integer(self):
        self.assertEqual(price_check_tool.parse_price_number("16.990.000₫"), 16_990_000)
        self.assertEqual(price_check_tool.parse_price_number("16990000.0"), 16_990_000)

    def test_viettel_old_customer_voucher_reduces_eligible_iphone_price(self):
        target = price_check_tool.Target("iPhone 17 Pro 256GB", "Viettel", "https://example.test/iphone-17-pro")
        visible = "Giá 33.190.000đ Giảm thêm 500.000đ cho khách hàng cũ của Viettel Store"

        result = price_check_tool.final_price_result(target, 33_190_000, visible, "2026-07-17T00:00:00", "structured_offer_price")

        self.assertEqual(result.value, 32_690_000)
        self.assertEqual(result.base_value, 33_190_000)
        self.assertEqual(result.discount, 500_000)

    def test_viettel_student_discount_reduces_ipad_price(self):
        target = price_check_tool.Target("iPad Gen 11th 128GB Wifi", "Viettel", "https://example.test/ipad-a16")
        visible = "Giá 11.990.000đ Giảm thêm 200.000đ cho học sinh, sinh viên, giáo viên"

        result = price_check_tool.final_price_result(target, 11_990_000, visible, "2026-07-24T00:00:00", "structured_offer_price")

        self.assertEqual(result.value, 11_790_000)
        self.assertEqual(result.discount, 200_000)

    def test_viettel_bundle_or_tammi_discount_does_not_reduce_price(self):
        target = price_check_tool.Target("iPhone 17 Pro Max 256GB", "Viettel", "https://example.test/iphone-17-pro-max")
        visible = "Giá 35.990.000đ Giảm thêm 3.000.000đ cho khách hàng mua kèm gói cước Tammi"

        result = price_check_tool.final_price_result(target, 35_990_000, visible, "2026-07-24T00:00:00", "structured_offer_price")

        self.assertEqual(result.value, 35_990_000)
        self.assertEqual(result.discount, 0)

    def test_viettel_image_only_old_customer_campaign_reduces_eligible_model(self):
        target = price_check_tool.Target("iPhone 17 Pro Max 256GB", "Viettel", "https://example.test/iphone-17-pro-max")
        visible = "Giá 35.990.000đ Quà tặng và ưu đãi khác"

        result = price_check_tool.final_price_result(target, 35_990_000, visible, "2026-07-24T00:00:00", "structured_offer_price")

        self.assertEqual(result.value, 35_490_000)
        self.assertEqual(result.discount, 500_000)

    def test_product_oos_wins_over_visible_price(self):
        raw_html = (
            '<html><body>Giá tại 16.990.000₫ '
            '<div class="productstatus">Hàng sắp về Đăng ký nhận thông tin</div></body></html>'
        )
        target = price_check_tool.Target("iPhone 16e 128GB", "MW", "https://example.test/iphone-16e")

        result = price_check_tool.extract_price(raw_html, target)

        self.assertEqual(result.status, "OOS")
        self.assertEqual(result.value, "OOS")
        self.assertEqual(result.availability, "OOS")

    def test_fpt_contact_advisory_wins_over_variant_price(self):
        raw_html = (
            '<html><body><div class="variant-data">finalPrice:11990000</div>'
            '<div id="noPrice"><p>Liên hệ tư vấn</p></div></body></html>'
        )
        target = price_check_tool.Target(
            "iPhone 13 128GB", "FPT", "https://example.test/iphone-13"
        )

        result = price_check_tool.extract_price(raw_html, target)

        self.assertEqual(result.status, "OOS")
        self.assertEqual(result.value, "OOS")
        self.assertEqual(result.availability_method, "fpt_product_status_element")

    def test_shopdunk_conflicting_server_ctas_require_render_confirmation(self):
        raw_html = (
            '<div id="cart-button-prd">'
            '<button class="add-to-cart-button">Mua ngay</button>'
            '<button class="add-to-cart-subscriber d-none">Sắp về hàng</button>'
            '</div>'
        )

        action = price_check_tool.extract_purchase_action(raw_html)

        self.assertEqual(action[0], "UNVERIFIED")
        self.assertEqual(action[1], "html_action_conflict")

    def test_shopdunk_stock_gate_overrides_price_metadata(self):
        target = price_check_tool.Target(
            "iPhone 13 128GB", "Shopdunk", "https://shopdunk.com/iphone-13"
        )
        result = price_check_tool.Result(
            target.model,
            target.retailer,
            target.url,
            12_590_000,
            "OK",
            html_price=12_590_000,
            source_method="shopdunk_price_value",
        )
        with mock.patch.object(
            price_check_tool,
            "shopdunk_stock_probe",
            return_value=("OOS", "shopdunk_stock_endpoint", "sku=SKU; finalGate=false"),
        ):
            result = price_check_tool.apply_shopdunk_stock_probe(result, target)

        self.assertEqual(result.status, "OOS")
        self.assertEqual(result.value, "OOS")
        self.assertEqual(result.availability_method, "shopdunk_stock_endpoint")
        self.assertTrue(price_check_tool.is_trusted_oos_method(result.availability_method))

    def test_unrelated_oos_far_from_product_price_does_not_override(self):
        raw_html = (
            "<html><body>Giá tại 16.990.000₫ "
            + ("Chi tiết sản phẩm đang bán. " * 40)
            + "Sản phẩm liên quan tạm hết hàng</body></html>"
        )
        target = price_check_tool.Target("iPhone 16e 128GB", "MW", "https://example.test/iphone-16e")

        result = price_check_tool.extract_price(raw_html, target)

        self.assertEqual(result.status, "OK")
        self.assertEqual(result.value, 16_990_000)

    def test_mw_review_or_promo_oos_text_does_not_override_buyable_price(self):
        raw_html = (
            "<html><body>Giá tại 35.590.000₫ Thêm vào giỏ Mua ngay "
            "Khách hàng đánh giá: Hàng sắp về ở nơi khác</body></html>"
        )
        target = price_check_tool.Target("iPad Pro M5 11 256GB", "MW", "https://example.test/ipad-pro")

        result = price_check_tool.extract_price(raw_html, target)

        self.assertEqual(result.status, "OK")
        self.assertEqual(result.value, 35_590_000)

    def test_mw_explicit_discontinued_product_status_is_oos(self):
        raw_html = (
            '<html><body>Giá tại 16.990.000₫ '
            '<div class="productstatus">SẢN PHẨM Ngừng kinh doanh</div></body></html>'
        )
        target = price_check_tool.Target("iPhone 14 128GB", "MW", "https://example.test/iphone-14")

        result = price_check_tool.extract_price(raw_html, target)

        self.assertEqual(result.status, "OOS")
        self.assertEqual(result.availability_method, "product_status_element")

    def test_shopdunk_hidden_variant_oos_does_not_override_html_price(self):
        raw_html = '<html><body><span id="price-value-1">17.290.000₫</span> Out of stock</body></html>'
        target = price_check_tool.Target("iPhone 17e", "Shopdunk", "https://example.test/iphone-17e")

        result = price_check_tool.extract_price(raw_html, target)

        self.assertEqual(result.status, "OK")
        self.assertEqual(result.value, 17_290_000)


    def test_viettel_structured_price_is_rendered_for_stock_verification(self):
        result = price_check_tool.Result(
            "iPad Air M3 11 128GB Wifi",
            "Viettel",
            "https://example.test/ipad-air-m3",
            16_990_000,
            "OK",
            source_method="structured_offer_price",
            risk_flags="WARN_WEAK_SOURCE",
        )

        self.assertTrue(price_check_tool.is_render_candidate(result))

    def test_render_backcheck_batches_large_candidate_sets(self):
        results = [
            price_check_tool.Result(
                f"Model {index}",
                "Viettel",
                f"https://example.test/{index}",
                17_290_000,
                "OK",
                source_method="structured_offer_price",
                risk_flags="WARN_WEAK_SOURCE",
                previous_week_price=17_290_000,
            )
            for index in range(68)
        ]
        report_dir = Path(tempfile.mkdtemp())
        calls = []

        def fake_run(args, **kwargs):
            input_path = Path(args[2])
            tasks = json.loads(input_path.read_text(encoding="utf-8"))
            calls.append(len(tasks))
            if len(tasks) > 40:
                return subprocess.CompletedProcess(
                    args=args,
                    returncode=1,
                    stdout="",
                    stderr="browser crashed on oversized batch",
                )
            rendered = [
                {
                    "key": task["key"],
                    "url": task["url"],
                    "ok": True,
                    "price": 17_290_000,
                    "method": "rendered_dom:.product-detail .price",
                    "actionStatus": "IN_STOCK",
                    "actionMethod": "rendered_action:.product-detail button",
                    "actionText": "MUA NGAY",
                    "availability": "",
                    "availabilityMethod": "",
                    "availabilityText": "",
                }
                for task in tasks
            ]
            return subprocess.CompletedProcess(
                args=args,
                returncode=0,
                stdout=json.dumps(rendered),
                stderr="",
            )

        with mock.patch.object(price_check_tool.subprocess, "run", side_effect=fake_run):
            price_check_tool.run_render_backcheck(results, report_dir)

        self.assertEqual(calls, [40, 28])
        self.assertTrue(all(result.rendered_price == 17_290_000 for result in results))
        self.assertTrue(all("WARN_RENDER_UNAVAILABLE" not in result.risk_flags for result in results))
        self.assertTrue(all(result.confidence == "VERIFIED_RENDER" for result in results))

if __name__ == "__main__":
    unittest.main()
