import importlib.util
import csv
import sys
import tempfile
import unittest
import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill


MODULE_PATH = Path(__file__).parents[1] / "tools" / "daily_price_check.py"


def load_daily_module(test_case):
    if not MODULE_PATH.exists():
        test_case.fail("daily_price_check.py has not been implemented")
    spec = importlib.util.spec_from_file_location("daily_price_check", MODULE_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class DailyReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.daily = load_daily_module(self)
        self.price = self.daily.price_check

    def result(self, value, status="OK", **kwargs):
        return self.price.Result(
            model="iPhone 17e 256GB",
            retailer="FPT",
            url="https://example.test/iphone-17e",
            value=value,
            status=status,
            fetched_at="2026-09-10T11:01:00",
            confidence=kwargs.pop("confidence", "HTML_OK"),
            **kwargs,
        )

    def test_first_successful_observation_creates_baseline_without_change(self):
        reconciliation = self.daily.reconcile_results(
            {},
            [self.result(17_490_000)],
            run_id="daily-20260910-primary",
            observed_at="2026-09-10T11:01:00+07:00",
        )

        latest = reconciliation.latest["iPhone 17e 256GB||FPT"]
        self.assertEqual(latest["current_value"], 17_490_000)
        self.assertEqual(latest["first_seen"], "2026-09-10T11:01:00+07:00")
        self.assertEqual(reconciliation.changes, [])

    def test_price_drop_creates_one_pending_change_when_result_is_risky(self):
        previous = {
            "iPhone 17e 256GB||FPT": {
                "key": "iPhone 17e 256GB||FPT",
                "model": "iPhone 17e 256GB",
                "retailer": "FPT",
                "current_value": 17_490_000,
                "stock_status": "In stock",
                "first_seen": "2026-09-09T11:00:00+07:00",
                "last_changed": "2026-09-09T11:00:00+07:00",
            }
        }

        reconciliation = self.daily.reconcile_results(
            previous,
            [self.result(16_990_000, risk_flags="WARN_OUTLIER", confidence="REVIEW")],
            run_id="daily-20260910-primary",
            observed_at="2026-09-10T11:01:00+07:00",
        )

        change = reconciliation.changes[0]
        self.assertEqual(change["change_type"], "Price decrease")
        self.assertEqual(change["old_value"], 17_490_000)
        self.assertEqual(change["new_value"], 16_990_000)
        self.assertEqual(change["delta_vnd"], -500_000)
        self.assertEqual(change["review_status"], "Pending")

    def test_stock_transition_creates_change(self):
        previous = {
            "iPhone 17e 256GB||FPT": {
                "key": "iPhone 17e 256GB||FPT",
                "model": "iPhone 17e 256GB",
                "retailer": "FPT",
                "current_value": 17_490_000,
                "stock_status": "In stock",
                "first_seen": "2026-09-09T11:00:00+07:00",
                "last_changed": "2026-09-09T11:00:00+07:00",
            }
        }

        reconciliation = self.daily.reconcile_results(
            previous,
            [self.result("OOS", status="OOS", confidence="OOS_CONFIRMED")],
            run_id="daily-20260910-primary",
            observed_at="2026-09-10T11:01:00+07:00",
        )

        self.assertEqual(reconciliation.changes[0]["change_type"], "Stock: In stock to OOS")
        self.assertEqual(reconciliation.latest["iPhone 17e 256GB||FPT"]["stock_status"], "OOS")

    def test_fetch_failure_preserves_previous_latest_value(self):
        previous = {
            "iPhone 17e 256GB||FPT": {
                "key": "iPhone 17e 256GB||FPT",
                "model": "iPhone 17e 256GB",
                "retailer": "FPT",
                "current_value": 17_490_000,
                "stock_status": "In stock",
                "first_seen": "2026-09-09T11:00:00+07:00",
                "last_changed": "2026-09-09T11:00:00+07:00",
                "last_checked": "2026-09-09T11:00:00+07:00",
            }
        }
        failed = self.result(None, status="ERROR", error="timeout", confidence="REVIEW")

        reconciliation = self.daily.reconcile_results(
            previous,
            [failed],
            run_id="daily-20260910-primary",
            observed_at="2026-09-10T11:01:00+07:00",
        )

        latest = reconciliation.latest["iPhone 17e 256GB||FPT"]
        self.assertEqual(latest["current_value"], 17_490_000)
        self.assertEqual(latest["last_checked"], "2026-09-09T11:00:00+07:00")
        self.assertEqual(reconciliation.changes, [])
        self.assertEqual(len(reconciliation.failures), 1)

    def test_existing_event_id_prevents_duplicate_change(self):
        previous = {
            "iPhone 17e 256GB||FPT": {
                "key": "iPhone 17e 256GB||FPT",
                "model": "iPhone 17e 256GB",
                "retailer": "FPT",
                "current_value": 17_490_000,
                "stock_status": "In stock",
                "first_seen": "2026-09-09T11:00:00+07:00",
                "last_changed": "2026-09-09T11:00:00+07:00",
            }
        }
        result = self.result(16_990_000)
        first = self.daily.reconcile_results(
            previous,
            [result],
            run_id="daily-20260910-primary",
            observed_at="2026-09-10T11:01:00+07:00",
        )

        second = self.daily.reconcile_results(
            previous,
            [result],
            run_id="daily-20260910-primary",
            observed_at="2026-09-10T11:01:00+07:00",
            existing_event_ids={first.changes[0]["event_id"]},
        )

        self.assertEqual(second.changes, [])


class DailyWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.daily = load_daily_module(self)
        if not hasattr(self.daily, "LATEST_HEADERS"):
            self.fail("daily workbook storage has not been implemented")
        self.price = self.daily.price_check
        self.temp_dir = Path(tempfile.mkdtemp())
        self.workbook_path = self.temp_dir / "Price Check Daily.xlsx"
        workbook = Workbook()
        latest = workbook.active
        latest.title = "Latest"
        latest.append(self.daily.LATEST_HEADERS)
        changes = workbook.create_sheet("Changes")
        changes.append(self.daily.CHANGE_HEADERS)
        run_log = workbook.create_sheet("Run Log")
        run_log.append(self.daily.RUN_LOG_HEADERS)
        workbook.save(self.workbook_path)

    def result(self, value, status="OK", **kwargs):
        return self.price.Result(
            model="iPhone 17e 256GB",
            retailer="FPT",
            url="https://example.test/iphone-17e",
            value=value,
            status=status,
            fetched_at="2026-09-10T11:01:00",
            confidence=kwargs.pop("confidence", "HTML_OK"),
            **kwargs,
        )

    def metadata(self, run_id, observed_at):
        return self.daily.RunMetadata(
            run_id=run_id,
            parent_run_id="",
            run_type="primary",
            started_at=observed_at,
            completed_at=observed_at,
            snapshot_path="outputs/daily/snapshots/sample.csv",
        )

    def test_workbook_records_baseline_then_risky_change_in_yellow(self):
        first_at = "2026-09-10T11:01:00+07:00"
        self.daily.apply_run_to_workbook(
            self.workbook_path,
            [self.result(17_490_000)],
            self.metadata("daily-20260910-primary", first_at),
            self.temp_dir / "backups",
        )
        second_at = "2026-09-11T11:01:00+07:00"
        summary = self.daily.apply_run_to_workbook(
            self.workbook_path,
            [self.result(16_990_000, risk_flags="WARN_OUTLIER", confidence="REVIEW")],
            self.metadata("daily-20260911-primary", second_at),
            self.temp_dir / "backups",
        )

        workbook = load_workbook(self.workbook_path, data_only=False)
        self.assertEqual(workbook["Latest"].max_row, 2)
        self.assertEqual(workbook["Changes"].max_row, 2)
        self.assertEqual(workbook["Run Log"].max_row, 3)
        self.assertEqual(summary.change_count, 1)
        review_col = self.daily.CHANGE_HEADERS.index("Review Status") + 1
        self.assertEqual(workbook["Changes"].cell(2, review_col).value, "Pending")
        self.assertEqual(workbook["Changes"].cell(2, 1).fill.fgColor.rgb, "00FFF2CC")
        self.assertTrue(any((self.temp_dir / "backups").glob("Price Check Daily_backup_*.xlsx")))

    def test_duplicate_run_id_does_not_append_rows(self):
        observed_at = "2026-09-10T11:01:00+07:00"
        metadata = self.metadata("daily-20260910-primary", observed_at)
        first = self.daily.apply_run_to_workbook(
            self.workbook_path,
            [self.result(17_490_000)],
            metadata,
            self.temp_dir / "backups",
        )
        second = self.daily.apply_run_to_workbook(
            self.workbook_path,
            [self.result(17_490_000)],
            metadata,
            self.temp_dir / "backups",
        )

        workbook = load_workbook(self.workbook_path)
        self.assertFalse(first.already_recorded)
        self.assertTrue(second.already_recorded)
        self.assertEqual(workbook["Run Log"].max_row, 2)

    def test_retry_target_reader_returns_only_failed_rows(self):
        snapshot = self.temp_dir / "price_snapshot_20260910-primary.csv"
        with snapshot.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=["model", "retailer", "url", "status"])
            writer.writeheader()
            writer.writerow({"model": "iPhone A", "retailer": "FPT", "url": "https://example.test/a", "status": "ERROR"})
            writer.writerow({"model": "iPhone B", "retailer": "MW", "url": "https://example.test/b", "status": "OK"})

        targets = self.daily.retry_targets_from_snapshot(snapshot)

        self.assertEqual([(target.model, target.retailer) for target in targets], [("iPhone A", "FPT")])

    def test_formatted_empty_tail_does_not_push_logs_to_row_5001(self):
        workbook = load_workbook(self.workbook_path)
        workbook["Changes"]["A5000"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
        workbook["Run Log"]["A5000"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
        workbook.save(self.workbook_path)
        observed_at = "2026-09-10T11:01:00+07:00"

        self.daily.apply_run_to_workbook(
            self.workbook_path,
            [self.result(17_490_000)],
            self.metadata("daily-20260910-primary", observed_at),
            self.temp_dir / "backups",
        )

        compact = load_workbook(self.workbook_path)
        self.assertEqual(compact["Changes"].max_row, 1)
        self.assertEqual(compact["Run Log"].max_row, 2)


class DailyOrchestrationTests(unittest.TestCase):
    def setUp(self):
        self.daily = load_daily_module(self)
        if not hasattr(self.daily, "make_run_id"):
            self.fail("daily orchestration has not been implemented")
        self.price = self.daily.price_check
        self.temp_dir = Path(tempfile.mkdtemp())

    def test_primary_run_id_is_stable_for_the_scheduled_date(self):
        now = dt.datetime(2026, 9, 10, 11, 3, 4, tzinfo=dt.timezone(dt.timedelta(hours=7)))

        self.assertEqual(self.daily.make_run_id("primary", now), "daily-20260910-primary")
        self.assertEqual(self.daily.make_run_id("retry", now), "daily-20260910-retry")
        self.assertEqual(self.daily.make_run_id("manual", now), "daily-20260910-manual-110304")

    def test_snapshot_round_trip_preserves_failures_for_retry(self):
        results = [
            self.price.Result("iPhone A", "FPT", "https://example.test/a", None, "ERROR", error="timeout"),
            self.price.Result("iPhone B", "MW", "https://example.test/b", 18_000_000, "OK"),
        ]

        snapshot = self.daily.write_daily_snapshot(results, self.temp_dir, "daily-20260910-primary")
        retry_targets = self.daily.retry_targets_from_snapshot(snapshot)

        self.assertEqual(snapshot.name, "price_snapshot_daily-20260910-primary.csv")
        self.assertEqual([(target.model, target.retailer) for target in retry_targets], [("iPhone A", "FPT")])

    def test_daily_message_reports_zero_changes_and_retry_state(self):
        summary = self.daily.WorkbookSummary(
            run_id="daily-20260910-primary",
            total_targets=195,
            ok_count=150,
            oos_count=43,
            error_count=2,
            change_count=0,
            review_count=3,
            backup_path="outputs/daily/backups/sample.xlsx",
        )

        message = self.daily.daily_summary_message(summary, [], retry_at="11:30")

        self.assertIn("0 biến động", message)
        self.assertIn("2 lỗi", message)
        self.assertIn("retry lúc 11:30", message)
        self.assertIn("3 cần kiểm tra thủ công", message)


if __name__ == "__main__":
    unittest.main()
